from __future__ import annotations

import json
import re
import sys
from dataclasses import asdict, dataclass
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pypdf import PdfReader

from config import OUTPUT_DIR, PROJECT_DIR, REPORT_DIR, SUBMISSION_STEM, TITLE, WEB_DIR


@dataclass
class Check:
    category: str
    name: str
    passed: bool
    detail: str


checks: list[Check] = []


def record(category: str, name: str, passed: bool, detail: str) -> None:
    checks.append(Check(category, name, bool(passed), detail))


def close_to(value: float, target: float, tolerance: float) -> bool:
    return abs(value - target) <= tolerance


def validate_html() -> None:
    index_path = WEB_DIR / "index.html"
    artifact_path = WEB_DIR / "artifact.json"
    receipt_path = WEB_DIR / "delivery_receipt.json"
    for path in (index_path, artifact_path, receipt_path):
        record("HTML", f"{path.name} 存在", path.exists() and path.stat().st_size > 0, str(path))
    if not index_path.exists():
        return
    html = index_path.read_text(encoding="utf-8")
    record("HTML", "正式标题", TITLE in html, TITLE)
    external_script = re.search(r'<script[^>]+src=["\']https?://', html, re.I)
    external_css = re.search(r'<link[^>]+href=["\']https?://', html, re.I)
    record("HTML", "无外部脚本依赖", external_script is None, "可离线打开")
    record("HTML", "无外部样式依赖", external_css is None, "可离线打开")
    if receipt_path.exists():
        receipt = json.loads(receipt_path.read_text(encoding="utf-8"))
        record("HTML", "官方打包器校验", bool(receipt.get("ok")), json.dumps(receipt, ensure_ascii=False))


def validate_docx() -> None:
    path = REPORT_DIR / f"{SUBMISSION_STEM}.docx"
    record("DOCX", "文件存在", path.exists() and path.stat().st_size > 0, str(path))
    if not path.exists():
        return
    document = Document(path)
    full_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    record("DOCX", "正式标题", TITLE in full_text, TITLE)
    record("DOCX", "5 幅统计图", len(document.inline_shapes) == 5, f"实际 {len(document.inline_shapes)} 幅")
    record("DOCX", "统计表完整", len(document.tables) >= 4, f"实际 {len(document.tables)} 个表格")
    figure_numbers = set(re.findall(r"^图\s*([1-5])", full_text, re.M))
    record("DOCX", "图号与标题", len(figure_numbers) == 5, f"识别到 {len(figure_numbers)} 个不同图号")

    section = document.sections[0]
    width_cm = section.page_width.cm
    height_cm = section.page_height.cm
    record(
        "DOCX",
        "A4 页面",
        close_to(width_cm, 21.0, 0.1) and close_to(height_cm, 29.7, 0.1),
        f"{width_cm:.2f} cm × {height_cm:.2f} cm",
    )

    normal = document.styles["Normal"]
    font_name = normal.font.name or ""
    font_size = normal.font.size.pt if normal.font.size else 0.0
    record("DOCX", "正文宋体", font_name in {"宋体", "SimSun"}, font_name)
    record("DOCX", "正文五号", close_to(font_size, 10.5, 0.05), f"{font_size:.1f} pt")
    fmt = normal.paragraph_format
    spacing = float(fmt.line_spacing) if isinstance(fmt.line_spacing, (int, float)) else 0.0
    before = fmt.space_before.pt if fmt.space_before else 0.0
    after = fmt.space_after.pt if fmt.space_after else 0.0
    record("DOCX", "1.5 倍行距", close_to(spacing, 1.5, 0.01), f"{spacing:g} 倍")
    record("DOCX", "0 磅段间距", close_to(before, 0.0, 0.01) and close_to(after, 0.0, 0.01), f"段前 {before:g} / 段后 {after:g}")
    record(
        "DOCX",
        "正文两端对齐",
        fmt.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY,
        str(fmt.alignment),
    )
    record("DOCX", "无模板占位符", "[[" not in full_text and "]]" not in full_text, "未发现 [[...]]")

    with ZipFile(path) as package:
        document_xml = package.read("word/document.xml").decode("utf-8")
    alt_count = len(re.findall(r'<wp:docPr[^>]+(?:descr|title)="[^"]+"', document_xml))
    record("DOCX", "图片替代文本", alt_count >= 5, f"识别到 {alt_count} 幅带说明的图片")


def validate_pdf() -> None:
    path = REPORT_DIR / f"{SUBMISSION_STEM}.pdf"
    record("PDF", "文件存在", path.exists() and path.stat().st_size > 0, str(path))
    if not path.exists():
        return
    reader = PdfReader(path)
    record("PDF", "页数合理", 6 <= len(reader.pages) <= 12, f"实际 {len(reader.pages)} 页")
    all_a4 = True
    sizes: list[str] = []
    for page in reader.pages:
        width = float(page.mediabox.width)
        height = float(page.mediabox.height)
        sizes.append(f"{width:.1f}×{height:.1f}")
        all_a4 &= close_to(width, 595.3, 2.0) and close_to(height, 841.9, 2.0)
    record("PDF", "全部页面为 A4", all_a4, ", ".join(sorted(set(sizes))))
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    compact_text = re.sub(r"\s+", "", text)
    compact_title = re.sub(r"\s+", "", TITLE)
    record("PDF", "正式标题可检索", compact_title in compact_text, TITLE)
    core_terms = all(term in text for term in ("逻辑回归", "决策树", "随机森林", "混淆矩阵", "AUC", "ROC"))
    record("PDF", "核心任务内容完整", core_terms, "三种算法与三类评价概念")
    figure_captions = re.findall(r"图\s*[1-5]", text)
    record("PDF", "图号可检索", len(set(figure_captions)) >= 5, f"识别到 {len(set(figure_captions))} 个不同图号")


def write_report() -> None:
    passed = sum(check.passed for check in checks)
    total = len(checks)
    payload = {
        "status": "PASS" if passed == total else "FAIL",
        "passed": passed,
        "total": total,
        "checks": [asdict(check) for check in checks],
        "submission_note": "提交前请把文件名中的“姓名”替换为真实姓名。",
    }
    (OUTPUT_DIR / "artifact_validation.json").write_text(
        json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    lines = [
        "# TASK5 最终交付校验报告",
        "",
        f"结论：**{payload['status']}（{passed}/{total}）**",
        "",
        "| 类别 | 校验项 | 结果 | 说明 |",
        "|---|---|---:|---|",
    ]
    for check in checks:
        detail = check.detail.replace("|", "\\|").replace("\n", " ")
        lines.append(f"| {check.category} | {check.name} | {'通过' if check.passed else '未通过'} | {detail} |")
    lines.extend(["", "> 提交前请把文件名中的“姓名”替换为真实姓名。", ""])
    (OUTPUT_DIR / "artifact_validation_report.md").write_text("\n".join(lines), encoding="utf-8")
    print(f"最终交付校验：{payload['status']}（{passed}/{total}）")
    if passed != total:
        for check in checks:
            if not check.passed:
                print(f"[FAIL] {check.category} / {check.name}: {check.detail}")
        raise SystemExit(1)


def main() -> int:
    validate_html()
    validate_docx()
    validate_pdf()
    write_report()
    return 0


if __name__ == "__main__":
    sys.exit(main())
