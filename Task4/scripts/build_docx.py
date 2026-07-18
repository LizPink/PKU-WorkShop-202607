from __future__ import annotations

import csv
from collections import defaultdict
from datetime import date
from pathlib import Path
from statistics import median

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


TASK_DIR = Path(__file__).resolve().parents[1]
OUTPUT_DIR = TASK_DIR / "outputs"
FINAL_DOCX = TASK_DIR / "李子平-Task4-海龟策略完整版.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
MUTED = "64748B"
LIGHT_FILL = "F2F4F7"
CALLOUT_FILL = "F4F6F9"
GOLD = "7A5A00"
RISK = "9B1C1C"
WHITE = "FFFFFF"


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def set_run_font(run, *, ascii_font: str = "Calibri", east_asia: str = "Microsoft YaHei", size: float | None = None, color: str | None = None, bold: bool | None = None, italic: bool | None = None) -> None:
    run.font.name = ascii_font
    run._element.get_or_add_rPr()
    run._element.rPr.rFonts.set(qn("w:ascii"), ascii_font)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), ascii_font)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_style_font(style, *, size: float, color: str, bold: bool = False, ascii_font: str = "Calibri", east_asia: str = "Microsoft YaHei") -> None:
    style.font.name = ascii_font
    style.font.size = Pt(size)
    style.font.color.rgb = rgb(color)
    style.font.bold = bold
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), ascii_font)
    rfonts.set(qn("w:hAnsi"), ascii_font)
    rfonts.set(qn("w:eastAsia"), east_asia)


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    set_style_font(normal, size=11, color="222222")
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    h1 = doc.styles["Heading 1"]
    set_style_font(h1, size=16, color=BLUE, bold=True)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(8)
    h1.paragraph_format.keep_with_next = True

    h2 = doc.styles["Heading 2"]
    set_style_font(h2, size=13, color=BLUE, bold=True)
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(6)
    h2.paragraph_format.keep_with_next = True

    h3 = doc.styles["Heading 3"]
    set_style_font(h3, size=12, color=DARK_BLUE, bold=True)
    h3.paragraph_format.space_before = Pt(8)
    h3.paragraph_format.space_after = Pt(4)
    h3.paragraph_format.keep_with_next = True

    caption = doc.styles["Caption"]
    set_style_font(caption, size=9.5, color=MUTED)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(10)

    for style_name, after in (("List Bullet", 8), ("List Number", 8)):
        style = doc.styles[style_name]
        set_style_font(style, size=11, color="222222")
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.167


def configure_page(section) -> None:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True


def add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=MUTED)
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_end)
    run2 = paragraph.add_run(" 页")
    set_run_font(run2, size=9, color=MUTED)


def configure_header_footer(section) -> None:
    header = section.header
    p = header.paragraphs[0]
    p.clear()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5))
    left = p.add_run("PKU Workshop 202607 · Task4")
    set_run_font(left, size=9, color=MUTED, bold=True)
    right = p.add_run("\t海龟交易策略")
    set_run_font(right, size=9, color=MUTED)

    first_header = section.first_page_header
    first_header.paragraphs[0].clear()

    footer = section.footer
    footer_p = footer.paragraphs[0]
    footer_p.clear()
    add_page_field(footer_p)
    first_footer = section.first_page_footer
    first_footer.paragraphs[0].clear()


def add_numbering(doc: Document, *, bullet: bool) -> int:
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(element.get(qn("w:abstractNumId"))) for element in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(element.get(qn("w:numId"))) for element in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet" if bullet else "decimal")
    level.append(num_fmt)
    text = OxmlElement("w:lvlText")
    text.set(qn("w:val"), "•" if bullet else "%1.")
    level.append(text)
    justification = OxmlElement("w:lvlJc")
    justification.set(qn("w:val"), "left")
    level.append(justification)
    ppr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    ppr.append(tabs)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "720")
    indent.set(qn("w:hanging"), "360")
    ppr.append(indent)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "160")
    spacing.set(qn("w:line"), "280")
    spacing.set(qn("w:lineRule"), "auto")
    ppr.append(spacing)
    level.append(ppr)
    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def add_list_item(doc: Document, text: str, num_id: int) -> None:
    paragraph = doc.add_paragraph(style="Normal")
    ppr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_element = OxmlElement("w:numId")
    num_id_element.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_element)
    ppr.append(num_pr)
    run = paragraph.add_run(text)
    set_run_font(run, size=11, color="222222")


def add_paragraph(doc: Document, text: str, *, bold_prefix: str | None = None) -> None:
    paragraph = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        bold = paragraph.add_run(bold_prefix)
        set_run_font(bold, size=11, color="222222", bold=True)
        rest = paragraph.add_run(text[len(bold_prefix):])
        set_run_font(rest, size=11, color="222222")
    else:
        run = paragraph.add_run(text)
        set_run_font(run, size=11, color="222222")


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int]) -> None:
    if sum(widths) != 9360:
        raise ValueError(f"表格列宽总和必须为9360 DXA，实际为{sum(widths)}")
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)
    for row in table.rows:
        row._tr.get_or_add_trPr()
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), "D7DEE8")


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[int], *, font_size: float = 9.5) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, text in enumerate(headers):
        cell = header.cells[index]
        shade_cell(cell, LIGHT_FILL)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(1)
        paragraph.paragraph_format.space_after = Pt(1)
        run = paragraph.add_run(text)
        set_run_font(run, size=font_size, color=INK, bold=True)
    for row_values in rows:
        row = table.add_row()
        for index, text in enumerate(row_values):
            cell = row.cells[index]
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if index == 0 else WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(1)
            paragraph.paragraph_format.space_after = Pt(1)
            paragraph.paragraph_format.line_spacing = 1.05
            run = paragraph.add_run(str(text))
            set_run_font(run, size=font_size, color="222222")
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(2)


def add_figure(doc: Document, image_path: Path, caption_text: str, width: float = 6.35) -> None:
    if not image_path.exists():
        raise FileNotFoundError(image_path)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    shape = run.add_picture(str(image_path), width=Inches(width))
    shape._inline.docPr.set("descr", caption_text)
    caption = doc.add_paragraph(style="Caption")
    caption.add_run(caption_text)


def add_formula(doc: Document, formula: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run(formula)
    set_run_font(run, ascii_font="Consolas", east_asia="Microsoft YaHei", size=10.5, color=INK)


def read_summary() -> list[dict[str, object]]:
    with (OUTPUT_DIR / "parameter_comparison.csv").open("r", encoding="utf-8-sig", newline="") as handle:
        rows: list[dict[str, object]] = []
        for row in csv.DictReader(handle):
            converted: dict[str, object] = dict(row)
            for field in (
                "entry_window", "exit_window", "atr_window", "observations", "buy_count", "sell_count",
                "completed_trades", "atr_stop_exits", "channel_exits", "open_position_at_end",
            ):
                converted[field] = int(float(row[field]))
            for field in (
                "stop_atr", "cumulative_return", "annualized_return", "annualized_volatility", "sharpe_ratio",
                "max_drawdown", "benchmark_return", "excess_return", "win_rate", "average_holding_days",
                "holding_ratio", "average_exposure", "total_transaction_cost", "total_turnover",
            ):
                try:
                    converted[field] = float(row[field])
                except ValueError:
                    converted[field] = float("nan")
            rows.append(converted)
    return rows


def pct(value: float) -> str:
    return "—" if value != value else f"{value:.1%}"


def num(value: float) -> str:
    return "—" if value != value else f"{value:.2f}"


def build_docx() -> Path:
    rows = read_summary()
    default_label = "E20/X10/ATR20/S2"
    default_oos = sorted(
        [row for row in rows if row["parameter"] == default_label and row["sample_period"] == "out_of_sample"],
        key=lambda row: float(row["sharpe_ratio"]),
        reverse=True,
    )
    positive_count = sum(float(row["cumulative_return"]) > 0 for row in default_oos)
    median_return = median(float(row["cumulative_return"]) for row in default_oos)
    median_sharpe = median(float(row["sharpe_ratio"]) for row in default_oos)
    median_mdd = median(float(row["max_drawdown"]) for row in default_oos)
    best_stock = default_oos[0]

    grouped: dict[str, list[dict[str, object]]] = defaultdict(list)
    for row in rows:
        if row["sample_period"] == "out_of_sample":
            grouped[str(row["parameter"])].append(row)
    parameter_rows: list[dict[str, object]] = []
    for parameter, group in grouped.items():
        parameter_rows.append(
            {
                "parameter": parameter,
                "median_return": median(float(row["cumulative_return"]) for row in group),
                "median_sharpe": median(float(row["sharpe_ratio"]) for row in group),
                "median_mdd": median(float(row["max_drawdown"]) for row in group),
                "positive_share": sum(float(row["cumulative_return"]) > 0 for row in group) / len(group),
            }
        )
    parameter_rows.sort(key=lambda row: float(row["median_sharpe"]), reverse=True)
    best_parameter = parameter_rows[0]

    doc = Document()
    configure_styles(doc)
    for section in doc.sections:
        configure_page(section)
        configure_header_footer(section)
    bullet_id = add_numbering(doc, bullet=True)
    number_id = add_numbering(doc, bullet=False)

    properties = doc.core_properties
    properties.title = "TASK4 复刻传奇：海龟交易法则实战演练"
    properties.subject = "海龟策略、ATR、通道突破、止损和回测"
    properties.author = "李子平"
    properties.last_modified_by = "OpenAI Codex"
    properties.keywords = "量化交易, 海龟策略, ATR, 最大回撤, 夏普比率, Python"

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(105)
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(16)
    run = kicker.add_run("QUANTITATIVE TRADING WORKSHOP · TASK4")
    set_run_font(run, size=10.5, color=GOLD, bold=True)
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_after = Pt(8)
    title_run = title_p.add_run("复刻传奇：海龟交易法则实战演练")
    set_run_font(title_run, size=28, color=INK, bold=True)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(26)
    subtitle_run = subtitle.add_run("高低点通道 · Wilder ATR · 风险定仓 · 止损 · 回测与参数研究")
    set_run_font(subtitle_run, size=13, color=DARK_BLUE)
    author = doc.add_paragraph()
    author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    author.paragraph_format.space_after = Pt(3)
    set_run_font(author.add_run("李子平"), size=12, color=INK, bold=True)
    prepared = doc.add_paragraph()
    prepared.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(prepared.add_run("北京大学工作坊 · 数据截至 2026-07-10"), size=10, color=MUTED)
    doc.add_page_break()

    doc.add_heading("技术摘要", level=1)
    add_list_item(doc, f"默认参数 {default_label} 在2024年以后的10只股票中有 {positive_count}/10 取得正累计回报；中位累计回报 {pct(median_return)}、中位夏普 {num(median_sharpe)}、中位最大回撤 {pct(median_mdd)}。", bullet_id)
    add_list_item(doc, f"默认参数样本外表现最好的是{best_stock['stock_name']}：累计回报 {pct(float(best_stock['cumulative_return']))}、最大回撤 {pct(float(best_stock['max_drawdown']))}、夏普 {num(float(best_stock['sharpe_ratio']))}。", bullet_id)
    add_list_item(doc, f"跨股票中位夏普最高的参数是 {best_parameter['parameter']}（中位夏普 {num(float(best_parameter['median_sharpe']))}）。参数排名只用于稳健性观察，不代表未来最优。", bullet_id)
    add_paragraph(doc, "结论：海龟策略是以规则化方式捕捉趋势、控制单笔风险的框架，不是高胜率预测模型。结果对股票、行情阶段和参数敏感，更适合多标的分散和样本外检验。", bold_prefix="结论：")

    doc.add_heading("1. 海龟策略的核心思想", level=1)
    add_paragraph(doc, "海龟交易法则是一类趋势跟踪系统。它不预测价格顶部或底部，而是等待市场突破近期高点后跟随趋势，用较短周期低点或波动率止损退出。策略允许多次小额亏损，以换取少数持续趋势带来的较大盈利。")
    doc.add_heading("1.1 关键优势", level=2)
    for item in (
        "规则客观：入场、退出、仓位和止损都可以写成明确公式，便于复现和审计。",
        "让利润奔跑：低点通道退出避免过早预测顶部，使持续趋势有机会积累收益。",
        "波动率统一：ATR把不同价格和波动水平的股票转换成风险单位。",
        "风险前置：交易前即可确定止损距离和最大计划损失。",
        "适合组合化：同一套规则可应用于多股票、多行业和多市场。",
    ):
        add_list_item(doc, item, bullet_id)
    doc.add_heading("1.2 主要局限", level=2)
    for item in (
        "横盘震荡会导致突破后迅速回落，连续产生小额止损。",
        "策略胜率通常不高，投资者需要承受较长的回撤和空仓期。",
        "跳空、涨跌停和流动性不足可能使实际成交价差于理论止损价。",
        "单只股票结果不稳定，参数优化容易产生数据挖掘偏差。",
    ):
        add_list_item(doc, item, bullet_id)

    doc.add_heading("2. 高低点通道、ATR与止损", level=1)
    doc.add_heading("2.1 高低点通道", level=2)
    add_paragraph(doc, "入场高点通道是过去N个交易日最高价的最大值；离场低点通道是过去M个交易日最低价的最小值。为避免使用当日尚未完成的信息，本实现将滚动通道整体向后移动1日。")
    add_formula(doc, "EntryHigh_t = max(High_{t-N}, …, High_{t-1})")
    add_formula(doc, "ExitLow_t = min(Low_{t-M}, …, Low_{t-1})")
    add_paragraph(doc, "默认N=20、M=10。收盘价首次突破入场通道后，下一交易日开盘买入；持仓后若收盘价跌破离场通道，则下一交易日开盘卖出。")

    doc.add_heading("2.2 真实波幅与Wilder ATR", level=2)
    add_paragraph(doc, "真实波幅TR不仅考虑当日最高价和最低价，还考虑相对昨收的向上或向下跳空。ATR是TR的平滑平均，用于表示当前正常波动尺度。")
    add_formula(doc, "TR_t = max(High_t − Low_t, |High_t − Close_{t−1}|, |Low_t − Close_{t−1}|)")
    add_formula(doc, "ATR_t = [(n−1) × ATR_{t−1} + TR_t] / n")
    add_paragraph(doc, "本实现使用20日Wilder ATR：第一个ATR取前20个TR的简单平均，之后按上式递推。")

    doc.add_heading("2.3 止损与风险定仓", level=2)
    add_paragraph(doc, "买入时设置初始止损价：入场价减去2倍信号日ATR。若交易日最低价触及止损线，则按止损价卖出；如果跳空低开已经低于止损价，则按开盘价成交。")
    add_formula(doc, "Stop = EntryPrice − k × ATR_entry")
    add_formula(doc, "Shares = min[Equity × 1% / (k × ATR_entry), Cash / EntryPrice]")
    add_paragraph(doc, "风险定仓使高波动股票自动减少股数。与完全满仓相比，它降低了单只股票对账户净值的影响，也使不同股票的回测更可比较。")

    doc.add_heading("2.4 默认参数与作用", level=2)
    add_table(
        doc,
        ["模块", "默认值", "含义"],
        [
            ["入场通道", "20日", "识别中期向上突破"],
            ["离场通道", "10日", "趋势转弱时退出"],
            ["ATR周期", "20日", "估计正常波动尺度"],
            ["止损倍数", "2ATR", "限制初始单笔风险"],
            ["账户风险", "1%", "决定股数而非固定满仓"],
            ["交易成本", "单边0.1%", "买卖时从现金扣除"],
        ],
        [1800, 1600, 5960],
    )

    doc.add_heading("3. Python实现与回测方法", level=1)
    add_paragraph(doc, "程序读取 Task3 已保存的10只股票前复权日线，统一日期和数值类型，并检查交易日重复、关键字段缺失、非正价格以及OHLC内部关系。随后计算通道、TR和ATR，再使用逐日现金—股数台账模拟买卖。")
    for step in (
        "加载并验证本地OHLCV行情。",
        "使用shift(1)计算20日最高通道和10日最低通道。",
        "计算TR和20日Wilder ATR。",
        "识别首次收盘突破，并在下一交易日开盘按ATR风险定仓买入。",
        "按低点通道或2ATR初始止损退出，同时记录成交价、成本和退出原因。",
        "逐日计算现金、股数、风险暴露、策略净值、基准净值和回撤。",
        "按全样本、样本内和样本外区间计算绩效并比较参数。",
    ):
        add_list_item(doc, step, number_id)

    doc.add_heading("3.1 绩效指标", level=2)
    add_formula(doc, "Cumulative Return = ∏(1 + r_t) − 1")
    add_formula(doc, "MDD = min(Equity_t / RunningMax_t − 1)")
    add_formula(doc, "Sharpe = mean(r_t) / std(r_t) × √252")
    add_paragraph(doc, "除累计回报、MDD和夏普外，报告还计算年化收益/波动、买入持有基准、超额收益、胜率、完成交易数、平均持有天数、持仓比例、平均资金暴露、总交易成本和退出类型。")

    doc.add_heading("4. 默认参数回测结果", level=1)
    add_paragraph(doc, f"样本外从2024-01-01开始。默认参数有{positive_count}只股票实现正累计回报，但相对买入持有基准普遍较低，主要原因是1%风险定仓使平均资金暴露较低。该设计强调风险可比性，而非最大化单边牛市收益。")
    result_table_rows = [
        [
            str(row["stock_name"]),
            str(row["industry"]),
            pct(float(row["cumulative_return"])),
            pct(float(row["benchmark_return"])),
            pct(float(row["max_drawdown"])),
            num(float(row["sharpe_ratio"])),
            str(int(row["completed_trades"])),
        ]
        for row in default_oos
    ]
    add_table(
        doc,
        ["股票", "行业", "累计回报", "基准回报", "最大回撤", "夏普", "交易数"],
        result_table_rows,
        [1400, 1300, 1450, 1450, 1300, 1100, 1360],
        font_size=8.5,
    )
    add_figure(doc, OUTPUT_DIR / "default_parameter_returns.png", "图1 默认参数下各股票样本外累计回报（2024-01至2026-07）")
    add_figure(doc, OUTPUT_DIR / "risk_return_scatter.png", "图2 默认参数样本外累计回报与最大回撤的联合比较")

    best_code = str(best_stock["ts_code"]).replace(".", "_")
    strategy_image = OUTPUT_DIR / f"{best_code}_T20_10_ATR20_S2p0_strategy.png"
    add_figure(doc, strategy_image, f"图3 {best_stock['stock_name']}的价格通道、ATR、交易信号、净值与回撤")

    doc.add_heading("5. 参数敏感性与行业差异", level=1)
    add_paragraph(doc, "参数研究同时改变入场/离场通道、ATR周期和止损倍数。下表按10只股票样本外中位夏普排序。E40/X20/ATR20/S2和E10/X5/ATR14/S1.5非常接近，但两者对交易频率和噪声的敏感度不同。")
    parameter_table_rows = [
        [
            str(row["parameter"]),
            pct(float(row["median_return"])),
            num(float(row["median_sharpe"])),
            pct(float(row["median_mdd"])),
            pct(float(row["positive_share"])),
        ]
        for row in parameter_rows
    ]
    add_table(
        doc,
        ["参数组合", "中位回报", "中位夏普", "中位MDD", "正收益占比"],
        parameter_table_rows,
        [2700, 1600, 1600, 1600, 1860],
        font_size=9,
    )
    add_figure(doc, OUTPUT_DIR / "stock_parameter_sharpe_heatmap.png", "图4 各股票在不同参数下的样本外夏普比率")
    add_figure(doc, OUTPUT_DIR / "industry_parameter_sharpe_heatmap.png", "图5 各行业在不同参数下的样本外中位夏普比率")

    doc.add_heading("5.1 参数观察", level=2)
    for item in (
        "短通道响应快、交易更频繁，能较早进入新趋势，但更容易被日常噪声触发。",
        "长通道减少无效突破，却可能延后入场，并在趋势反转后较晚退出。",
        "较小ATR倍数提高仓位但更易触发止损；较大倍数降低股数并扩大价格容忍区间。",
        "半导体样本在多数参数下表现较好，食品饮料样本持续偏弱；但每行业仅2只股票，不能外推到完整行业。",
        "同一参数在不同股票上结果差异显著，说明分散化和样本外验证比单点最优化更重要。",
    ):
        add_list_item(doc, item, bullet_id)

    doc.add_heading("6. 适应场景与使用心得", level=1)
    add_paragraph(doc, "海龟法则更适合存在持续趋势、流动性较好、可以跨标的分散的市场。它在单边上涨或下跌趋势中有机会保留盈利头寸，在窄幅震荡、消息跳空和成交受限时则容易失效。")
    for item in (
        "先确定风险预算，再讨论收益；ATR应真正进入止损和仓位公式，而不只是附加指标。",
        "通道必须使用前一日及更早数据，普通信号应错后一交易日执行。",
        "只看累计收益会掩盖风险，至少同时查看MDD、夏普、交易次数、胜率、平均暴露和基准收益。",
        "胜率低不代表策略无效；趋势策略通常依赖少数大盈利覆盖多次小亏损。",
        "参数应选跨股票、跨行业、样本外不过度失真的区域，不应只选全样本最高收益。",
    ):
        add_list_item(doc, item, bullet_id)

    doc.add_heading("7. 局限与下一步", level=1)
    for item in (
        "股票池为当前人工挑选的10只代表性股票，存在幸存者偏差。",
        "模型允许小数股，未模拟A股100股整手、涨跌停无法成交、停牌、滑点和容量限制。",
        "买卖单边0.1%成本、零无风险利率和前复权价格均为教学假设。",
        "样本外区间只有约2.5年，不能覆盖所有市场状态。",
    ):
        add_list_item(doc, item, bullet_id)
    add_paragraph(doc, "下一步应扩大到历史成分股，加入滚动样本外检验和组合级风险预算，并在相同成本和成交约束下与双均线及买入持有比较。")

    doc.add_heading("附录：成果与复现", level=1)
    add_paragraph(doc, "主要代码位于 Task4/scripts，回测明细、交易明细、汇总CSV和PNG位于 Task4/outputs；Task4_process_walkthrough.ipynb 为已执行的教学流程，Task4/web/index.html 为与Task1-Task3统一风格的静态网页报告。")
    code_p = doc.add_paragraph()
    code_p.paragraph_format.space_before = Pt(4)
    code_p.paragraph_format.space_after = Pt(8)
    code_p.paragraph_format.left_indent = Inches(0.25)
    code_run = code_p.add_run("uv --cache-dir .uv-cache run python .\\Task4\\scripts\\run_all.py")
    set_run_font(code_run, ascii_font="Consolas", east_asia="Microsoft YaHei", size=9.5, color=INK)
    add_paragraph(doc, "验证脚本独立复核通道错位、TR/ATR公式、资金恒等式、交易成本、退出原因、累计回报复利、MDD范围和交付物完整性。本报告用于课程教学与研究演示，不构成投资建议。")

    FINAL_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(FINAL_DOCX)
    return FINAL_DOCX


if __name__ == "__main__":
    print(build_docx())
