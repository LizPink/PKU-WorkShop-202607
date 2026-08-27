from __future__ import annotations

import json
import math
import re
import sys
from dataclasses import asdict, dataclass
from pathlib import Path
from zipfile import ZipFile

import numpy as np
import pandas as pd
from PIL import Image
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pypdf import PdfReader

from config import (
    FEATURE_COLUMNS,
    FIGURE_DIR,
    MODEL_ORDER,
    OUTPUT_DIR,
    PROCESSED_DIR,
    REPORT_DIR,
    SUBMISSION_STEM,
    TITLE,
    TOP_N,
    ensure_directories,
)


@dataclass
class Check:
    category: str
    name: str
    passed: bool
    detail: str
    severity: str = "High"


checks: list[Check] = []


def record(category: str, name: str, passed: bool, detail: str, severity: str = "High") -> None:
    checks.append(Check(category, name, bool(passed), str(detail), severity))


def close(left: float, right: float, tolerance: float = 1e-7) -> bool:
    return math.isclose(float(left), float(right), rel_tol=tolerance, abs_tol=tolerance)


def rank_ic(actual: pd.Series, predicted: pd.Series) -> float:
    return float(actual.rank().corr(predicted.rank()))


def validate_analysis() -> None:
    required = [
        PROCESSED_DIR / "quarterly_panel.csv",
        OUTPUT_DIR / "test_predictions.csv",
        OUTPUT_DIR / "model_metrics.csv",
        OUTPUT_DIR / "quarterly_ic.csv",
        OUTPUT_DIR / "quarterly_returns.csv",
        OUTPUT_DIR / "holdings.csv",
        OUTPUT_DIR / "backtest_metrics.csv",
        OUTPUT_DIR / "cost_sensitivity.csv",
        OUTPUT_DIR / "data_quality.json",
        OUTPUT_DIR / "selected_parameters.json",
    ]
    for path in required:
        record("Files", path.name, path.exists() and path.stat().st_size > 0, str(path))
    if not all(path.exists() for path in required):
        return

    panel = pd.read_csv(PROCESSED_DIR / "quarterly_panel.csv", dtype={"ts_code": str})
    predictions = pd.read_csv(OUTPUT_DIR / "test_predictions.csv", dtype={"ts_code": str})
    model_metrics = pd.read_csv(OUTPUT_DIR / "model_metrics.csv").set_index("model")
    quarterly_ic = pd.read_csv(OUTPUT_DIR / "quarterly_ic.csv")
    quarterly_returns = pd.read_csv(OUTPUT_DIR / "quarterly_returns.csv")
    holdings = pd.read_csv(OUTPUT_DIR / "holdings.csv", dtype={"ts_code": str})
    backtest_metrics = pd.read_csv(OUTPUT_DIR / "backtest_metrics.csv").set_index("model")
    quality = json.loads((OUTPUT_DIR / "data_quality.json").read_text(encoding="utf-8"))

    record("Data", "季度面板主键唯一", not panel.duplicated(["ts_code", "quarter_index"]).any(), f"duplicates={panel.duplicated(['ts_code','quarter_index']).sum()}")
    record("Data", "信号早于进场且进场早于退出", ((pd.to_datetime(panel["signal_date"]) < pd.to_datetime(panel["entry_date"])) & (pd.to_datetime(panel["entry_date"]) < pd.to_datetime(panel["exit_date"]))).all(), "signal < entry < exit")
    record("Data", "18 个模型因子完整", len(FEATURE_COLUMNS) == 18 and all(f"x_{feature}" in panel.columns for feature in FEATURE_COLUMNS), f"features={len(FEATURE_COLUMNS)}")
    record("Data", "数据质量元数据一致", quality["panel_rows"] == len(panel), f"metadata={quality['panel_rows']}, actual={len(panel)}")

    split = pd.read_csv(OUTPUT_DIR / "split_summary.csv").set_index("split")
    train_end = pd.Period(split.loc["train", "quarter_end"], freq="Q")
    validation_start = pd.Period(split.loc["validation", "quarter_start"], freq="Q")
    validation_end = pd.Period(split.loc["validation", "quarter_end"], freq="Q")
    test_start = pd.Period(split.loc["test", "quarter_start"], freq="Q")
    record("Method", "时间划分严格递增", train_end < validation_start <= validation_end < test_start, f"{train_end} < {validation_start} <= {validation_end} < {test_start}")
    record("Method", "测试集为 8 个季度", predictions["quarter_index"].nunique() == 8, f"quarters={predictions['quarter_index'].nunique()}")

    for model_key in MODEL_ORDER:
        pred_col = f"prediction_{model_key}"
        recomputed_ics = predictions.groupby("quarter_index").apply(lambda group: rank_ic(group["target_rank"], group[pred_col]), include_groups=False)
        saved_ics = quarterly_ic.loc[quarterly_ic["model"] == model_key].sort_values("quarter_index")["rank_ic"].to_numpy()
        record("Prediction", f"{model_key} 季度 Rank IC 独立复算", np.allclose(recomputed_ics.to_numpy(), saved_ics, atol=1e-8), f"max_diff={np.max(np.abs(recomputed_ics.to_numpy()-saved_ics)):.2e}")
        record("Prediction", f"{model_key} 平均 Rank IC 一致", close(recomputed_ics.mean(), model_metrics.loc[model_key, "mean_rank_ic"]), f"saved={model_metrics.loc[model_key,'mean_rank_ic']:.8f}, recomputed={recomputed_ics.mean():.8f}")

        model_holdings = holdings.loc[holdings["model"] == model_key]
        counts = model_holdings.groupby("quarter_index").size()
        weights = model_holdings.groupby("quarter_index")["weight"].sum()
        record("Portfolio", f"{model_key} 每季度 Top {TOP_N}", (counts == TOP_N).all(), counts.to_dict())
        record("Portfolio", f"{model_key} 权重和为 1", np.allclose(weights, 1.0), f"range={weights.min():.8f}..{weights.max():.8f}")
        recomputed_gross = model_holdings.groupby("quarter_index").apply(lambda group: np.average(group["forward_return"], weights=group["weight"]), include_groups=False)
        saved_quarterly = quarterly_returns.loc[quarterly_returns["model"] == model_key].sort_values("quarter_index").set_index("quarter_index")
        record("Portfolio", f"{model_key} Top 30 收益复算", np.allclose(recomputed_gross.loc[saved_quarterly.index], saved_quarterly["gross_return"]), f"max_diff={np.max(np.abs(recomputed_gross.loc[saved_quarterly.index]-saved_quarterly['gross_return'])):.2e}")
        record("Portfolio", f"{model_key} 净收益成本公式", np.allclose(saved_quarterly["net_return"], saved_quarterly["gross_return"] - saved_quarterly["trading_cost"]), "net = gross - cost")

        nav = (1 + saved_quarterly["net_return"]).cumprod()
        cumulative = nav.iloc[-1] - 1
        max_dd = (nav / nav.cummax() - 1).min()
        record("Performance", f"{model_key} 累计收益复算", close(cumulative, backtest_metrics.loc[model_key, "cumulative_return"]), f"saved={backtest_metrics.loc[model_key,'cumulative_return']:.8f}, recomputed={cumulative:.8f}")
        record("Performance", f"{model_key} 最大回撤复算", close(max_dd, backtest_metrics.loc[model_key, "max_drawdown"]), f"saved={backtest_metrics.loc[model_key,'max_drawdown']:.8f}, recomputed={max_dd:.8f}")


def validate_figures() -> None:
    expected = [
        "figure_1_sample_coverage.png",
        "figure_2_factor_correlation.png",
        "figure_3_quarterly_rank_ic.png",
        "figure_4_model_excess_return.png",
        "figure_5_cumulative_nav.png",
        "figure_6_quarterly_returns.png",
        "figure_7_drawdown.png",
        "figure_8_random_forest_feature_importance.png",
    ]
    for name in expected:
        path = FIGURE_DIR / name
        if path.exists():
            with Image.open(path) as image:
                width, height = image.size
            valid = width >= 1200 and height >= 800
            detail = f"{width}x{height}px, {path.stat().st_size} bytes"
        else:
            valid, detail = False, "missing"
        record("Figure", name, valid, detail)


def validate_docx() -> None:
    path = REPORT_DIR / f"{SUBMISSION_STEM}.docx"
    record("DOCX", "文件存在", path.exists() and path.stat().st_size > 0, str(path))
    if not path.exists():
        return
    document = Document(path)
    full_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    record("DOCX", "正式标题", TITLE in full_text, TITLE)
    record("DOCX", "8 幅统计图", len(document.inline_shapes) == 8, f"actual={len(document.inline_shapes)}")
    record("DOCX", "7 张核心表", len(document.tables) >= 7, f"actual={len(document.tables)}")
    figure_numbers = set(re.findall(r"^图\s*([1-8])", full_text, re.M))
    record("DOCX", "图号完整", len(figure_numbers) == 8, f"figure_numbers={sorted(figure_numbers)}")
    section = document.sections[0]
    record("DOCX", "A4 页面", abs(section.page_width.cm - 21.0) < 0.1 and abs(section.page_height.cm - 29.7) < 0.1, f"{section.page_width.cm:.2f}x{section.page_height.cm:.2f}cm")
    normal = document.styles["Normal"]
    font_name = normal.font.name or ""
    font_size = normal.font.size.pt if normal.font.size else 0
    spacing = float(normal.paragraph_format.line_spacing) if isinstance(normal.paragraph_format.line_spacing, (int, float)) else 0
    before = normal.paragraph_format.space_before.pt if normal.paragraph_format.space_before else 0
    after = normal.paragraph_format.space_after.pt if normal.paragraph_format.space_after else 0
    record("DOCX", "正文宋体", font_name in {"SimSun", "宋体"}, font_name)
    record("DOCX", "正文五号", abs(font_size - 10.5) < 0.05, f"{font_size:.1f}pt")
    record("DOCX", "正文 1.5 倍行距", abs(spacing - 1.5) < 0.01, str(spacing))
    record("DOCX", "正文 0 段间距", abs(before) < 0.01 and abs(after) < 0.01, f"before={before}, after={after}")
    record("DOCX", "正文两端对齐", normal.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY, str(normal.paragraph_format.alignment))
    record("DOCX", "无未替换占位符", "[[" not in full_text and "]]" not in full_text, "No [[...]] markers")
    with ZipFile(path) as package:
        document_xml = package.read("word/document.xml").decode("utf-8")
    alt_count = len(re.findall(r'<wp:docPr[^>]+(?:descr|title)="[^"]+"', document_xml))
    record("DOCX", "图片替代文本", alt_count >= 8, f"alt_count={alt_count}")


def validate_pdf() -> None:
    path = REPORT_DIR / f"{SUBMISSION_STEM}.pdf"
    record("PDF", "文件存在", path.exists() and path.stat().st_size > 0, str(path))
    if not path.exists():
        return
    reader = PdfReader(path)
    record("PDF", "页数合理", 12 <= len(reader.pages) <= 35, f"pages={len(reader.pages)}")
    a4 = all(abs(float(page.mediabox.width) - 595.3) < 2 and abs(float(page.mediabox.height) - 841.9) < 2 for page in reader.pages)
    record("PDF", "全部页面 A4", a4, f"pages={len(reader.pages)}")
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    compact = re.sub(r"\s+", "", text)
    record("PDF", "正式标题可检索", re.sub(r"\s+", "", TITLE) in compact, TITLE)
    terms = ("机器学习交易策略", "自变量", "应变量", "决策树", "随机森林", "Rank IC", "Top 30", "回测")
    record("PDF", "核心任务内容完整", all(term.replace(" ", "") in compact for term in terms), ", ".join(terms))
    figures = set(re.findall(r"图\s*([1-8])", text))
    record("PDF", "8 个图号可检索", len(figures) == 8, f"figures={sorted(figures)}")


def write_report() -> Path:
    failed = [check for check in checks if not check.passed]
    status = "Ready to share" if not failed else "Needs revision"
    payload = {"status": status, "passed": len(checks) - len(failed), "total": len(checks), "checks": [asdict(check) for check in checks]}
    (OUTPUT_DIR / "validation_checks.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    lines = [
        "# TASK6 Validation Report",
        "",
        f"## Overall Assessment: {status}",
        "",
        "### Methodology Review",
        "",
        "本校验覆盖季度主键、信号/交易时间顺序、时间划分、Rank IC、Top 30 持仓、权重、交易成本、累计收益、最大回撤、统计图、DOCX 排版和 PDF 内容。",
        "",
        "### Issues Found",
        "",
    ]
    if failed:
        for index, check in enumerate(failed, start=1):
            lines.append(f"{index}. [Severity: {check.severity}] {check.category}/{check.name}: {check.detail}")
    else:
        lines.append("未发现阻止提交的自动校验问题。")
    lines.extend(["", "### Calculation Spot-Checks", ""])
    for check in checks:
        lines.append(f"- [{'PASS' if check.passed else 'FAIL'}] {check.category}/{check.name}: {check.detail}")
    lines.extend(
        [
            "",
            "### Required Caveats for Readers",
            "",
            "- 当前中证 300 成分列表中等距抽取的 120 只股票被回溯使用，存在幸存者与样本选择偏差；基准是研究样本等权平均而非官方指数收益。",
            "- 历史 ST、停牌、涨跌停和冲击成本未完整模拟，成本为情景假设。",
            "- 测试期只有 8 个季度，结果不构成未来收益保证或投资建议。",
        ]
    )
    path = OUTPUT_DIR / "validation_report.md"
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def main() -> int:
    ensure_directories()
    validate_analysis()
    validate_figures()
    validate_docx()
    validate_pdf()
    report = write_report()
    failed = [check for check in checks if not check.passed]
    print(f"Validation: {len(checks)-len(failed)}/{len(checks)} passed -> {report}")
    if failed:
        for check in failed:
            print(f"[FAIL] {check.category}/{check.name}: {check.detail}")
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
