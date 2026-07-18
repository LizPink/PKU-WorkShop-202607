from __future__ import annotations

import json
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Twips

from config import (
    FEATURE_COLUMNS,
    FIGURE_DIR,
    MODEL_NAMES_ZH,
    MODEL_ORDER,
    ONE_WAY_COST,
    OUTPUT_DIR,
    REPORT_DIR,
    SUBMISSION_STEM,
    TITLE,
    TOP_N,
    ensure_directories,
)


FONT_BODY = "SimSun"
FONT_HEADING = "SimHei"
BODY_SIZE = 10.5
TABLE_SIZE = 9.0
CONTENT_WIDTH_DXA = 9025
INK = "263238"
BLUE = "2F6B8A"
GOLD = "C6912B"
MUTED = "6B747C"
LIGHT_BLUE = "E8F1F5"
LIGHT_GRAY = "F4F6F7"
GRID = "B7C1C8"


FEATURE_LABELS = {
    "ret_5d": ("动量/反转", "5 日收益率", "信号日前 5 个交易日复权收盘价收益"),
    "ret_21d": ("动量", "21 日收益率", "约 1 个月价格动量"),
    "ret_63d": ("动量", "63 日收益率", "约 1 个季度价格动量"),
    "ret_126d": ("动量", "126 日收益率", "约半年价格动量"),
    "ret_252d": ("动量", "252 日收益率", "约一年价格动量"),
    "momentum_12_1": ("动量", "12-1 月动量", "剔除最近 1 个月后的过去一年收益"),
    "volatility_20d": ("风险", "20 日年化波动率", "日收益标准差乘以年化系数"),
    "volatility_60d": ("风险", "60 日年化波动率", "较长窗口的收益波动"),
    "downside_volatility_60d": ("风险", "60 日下行波动率", "只保留下跌日收益的平方均值并年化"),
    "max_drawdown_126d": ("风险", "126 日最大回撤", "半年内价格相对滚动高点的最深跌幅"),
    "turnover_20d": ("流动性", "20 日平均换手率", "短期交易活跃度；备用数据源缺失时填中性值"),
    "turnover_60d": ("流动性", "60 日平均换手率", "中期交易活跃度；备用数据源缺失时填中性值"),
    "log_amount_20d": ("流动性", "20 日成交额对数", "成交规模的对数变换"),
    "amihud_20d": ("流动性", "20 日 Amihud 指标", "绝对收益除以成交量代理的滚动均值"),
    "volume_ratio_20_60": ("流动性", "20/60 日量比", "短期成交量相对中期成交量"),
    "price_to_ma20": ("趋势", "价格相对 MA20", "收盘价相对 20 日均线偏离"),
    "price_to_ma60": ("趋势", "价格相对 MA60", "收盘价相对 60 日均线偏离"),
    "intraday_range_20d": ("风险", "20 日日内振幅", "最高价与最低价差相对收盘价的滚动均值"),
}


def set_run_font(run, *, size: float = BODY_SIZE, bold: bool | None = None, color: str = INK, font: str = FONT_BODY, italic: bool | None = None) -> None:
    run.font.name = font
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:ascii"), font)
    rpr.rFonts.set(qn("w:hAnsi"), font)
    rpr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_paragraph_format(
    paragraph,
    *,
    alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
    line_spacing: float = 1.5,
    keep_with_next: bool = False,
    first_line_chars: float | None = 2,
    space_before: float = 0,
    space_after: float = 0,
) -> None:
    fmt = paragraph.paragraph_format
    fmt.alignment = alignment
    fmt.space_before = Pt(space_before)
    fmt.space_after = Pt(space_after)
    fmt.line_spacing = line_spacing
    fmt.keep_with_next = keep_with_next
    fmt.widow_control = True
    if first_line_chars is not None:
        fmt.first_line_indent = Pt(BODY_SIZE * first_line_chars)


def add_body(doc: Document, text: str, *, bold_lead: str | None = None) -> object:
    paragraph = doc.add_paragraph()
    set_paragraph_format(paragraph)
    if bold_lead and text.startswith(bold_lead):
        set_run_font(paragraph.add_run(bold_lead), bold=True)
        set_run_font(paragraph.add_run(text[len(bold_lead) :]))
    else:
        set_run_font(paragraph.add_run(text))
    return paragraph


def add_heading(doc: Document, text: str, level: int = 1) -> object:
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    set_paragraph_format(
        paragraph,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
        keep_with_next=True,
        first_line_chars=None,
        space_before=10 if level == 1 else 6,
        space_after=4,
    )
    paragraph.add_run(text)
    return paragraph


def add_bullet(doc: Document, text: str) -> object:
    paragraph = doc.add_paragraph(style="List Bullet")
    set_paragraph_format(paragraph, first_line_chars=None)
    paragraph.paragraph_format.left_indent = Cm(0.74)
    paragraph.paragraph_format.first_line_indent = Cm(-0.37)
    set_run_font(paragraph.add_run(text))
    return paragraph


def shade_paragraph(paragraph, fill: str) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    shd = ppr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        ppr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_callout(doc: Document, label: str, text: str) -> object:
    paragraph = doc.add_paragraph()
    set_paragraph_format(paragraph, alignment=WD_ALIGN_PARAGRAPH.LEFT, first_line_chars=None, space_before=4, space_after=4)
    shade_paragraph(paragraph, LIGHT_BLUE)
    set_run_font(paragraph.add_run(f"{label}："), bold=True, color=BLUE)
    set_run_font(paragraph.add_run(text))
    return paragraph


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def ensure_child(parent, tag: str):
    child = parent.find(qn(tag))
    if child is None:
        child = OxmlElement(tag)
        parent.append(child)
    return child


def set_cell_margins(cell, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = ensure_child(tc_pr, "w:tcMar")
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        margin = ensure_child(tc_mar, f"w:{side}")
        margin.set(qn("w:w"), str(value))
        margin.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int]) -> None:
    if sum(widths_dxa) != CONTENT_WIDTH_DXA:
        raise ValueError(f"Table widths must sum to {CONTENT_WIDTH_DXA}, got {sum(widths_dxa)}")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = ensure_child(tbl_pr, "w:tblW")
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = ensure_child(tbl_pr, "w:tblInd")
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    layout = ensure_child(tbl_pr, "w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for col_index, width in enumerate(widths_dxa):
        table.columns[col_index].width = Twips(width)
    for row in table.rows:
        row.height = None
        for col_index, cell in enumerate(row.cells):
            width = widths_dxa[col_index]
            cell.width = Twips(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = ensure_child(tc_pr, "w:tcW")
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def add_table_caption(doc: Document, text: str) -> object:
    paragraph = doc.add_paragraph(style="Caption")
    set_paragraph_format(paragraph, alignment=WD_ALIGN_PARAGRAPH.CENTER, keep_with_next=True, first_line_chars=None, space_before=4, space_after=4)
    set_run_font(paragraph.add_run(text), size=10, bold=True)
    return paragraph


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths_dxa: list[int], alignments: list[object] | None = None) -> object:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    repeat_table_header(table.rows[0])
    alignments = alignments or [WD_ALIGN_PARAGRAPH.LEFT] + [WD_ALIGN_PARAGRAPH.CENTER] * (len(headers) - 1)
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, LIGHT_BLUE)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        set_paragraph_format(paragraph, alignment=WD_ALIGN_PARAGRAPH.CENTER, first_line_chars=None, line_spacing=1.15)
        set_run_font(paragraph.add_run(header), size=TABLE_SIZE, bold=True)
    for row_data in rows:
        row = table.add_row()
        for index, value in enumerate(row_data):
            cell = row.cells[index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            set_paragraph_format(paragraph, alignment=alignments[index], first_line_chars=None, line_spacing=1.15)
            set_run_font(paragraph.add_run(str(value)), size=TABLE_SIZE)
    set_table_geometry(table, widths_dxa)
    return table


def add_figure(doc: Document, filename: str, caption: str, interpretation: str, width_cm: float = 15.2) -> None:
    image_path = FIGURE_DIR / filename
    if not image_path.exists():
        raise FileNotFoundError(image_path)
    paragraph = doc.add_paragraph()
    set_paragraph_format(paragraph, alignment=WD_ALIGN_PARAGRAPH.CENTER, keep_with_next=True, first_line_chars=None, space_before=4)
    inline_shape = paragraph.add_run().add_picture(str(image_path), width=Cm(width_cm))
    inline_shape._inline.docPr.set("descr", caption)
    inline_shape._inline.docPr.set("title", caption)
    caption_paragraph = doc.add_paragraph(style="Caption")
    set_paragraph_format(caption_paragraph, alignment=WD_ALIGN_PARAGRAPH.CENTER, keep_with_next=True, first_line_chars=None, space_after=2)
    set_run_font(caption_paragraph.add_run(caption), size=9.5, bold=True)
    add_body(doc, interpretation)


def add_page_number(paragraph) -> None:
    set_paragraph_format(paragraph, alignment=WD_ALIGN_PARAGRAPH.CENTER, first_line_chars=None, line_spacing=1.0)
    set_run_font(paragraph.add_run("第 "), size=9, color=MUTED)
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_run_font(run, size=9, color=MUTED)
    set_run_font(paragraph.add_run(" 页"), size=9, color=MUTED)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)
    section.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    normal.font.name = FONT_BODY
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT_BODY)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_BODY)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY)
    normal.font.size = Pt(BODY_SIZE)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    heading_tokens = {
        "Heading 1": (15, BLUE, 10, 4),
        "Heading 2": (12, INK, 6, 3),
        "Heading 3": (10.5, BLUE, 4, 2),
    }
    for style_name, (size, color, before, after) in heading_tokens.items():
        style = doc.styles[style_name]
        style.font.name = FONT_HEADING
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT_HEADING)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_HEADING)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_HEADING)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.name = FONT_BODY
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY)
        style.font.size = Pt(BODY_SIZE)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.line_spacing = 1.5

    caption = doc.styles["Caption"]
    caption.font.name = FONT_BODY
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY)
    caption.font.size = Pt(9.5)
    caption.font.color.rgb = RGBColor.from_string(INK)
    caption.paragraph_format.space_before = Pt(0)
    caption.paragraph_format.space_after = Pt(0)

    header = section.header
    hp = header.paragraphs[0]
    set_paragraph_format(hp, alignment=WD_ALIGN_PARAGRAPH.RIGHT, first_line_chars=None, line_spacing=1.0)
    set_run_font(hp.add_run("TASK6 · 智能决策者"), size=9, color=MUTED)
    add_page_number(section.footer.paragraphs[0])


def pct(value: float, digits: int = 2, signed: bool = False) -> str:
    sign = "+" if signed else ""
    return f"{value:{sign}.{digits}%}"


def num(value: float, digits: int = 3) -> str:
    return "-" if pd.isna(value) else f"{value:.{digits}f}"


def build_document() -> Path:
    ensure_directories()
    quality = json.loads((OUTPUT_DIR / "data_quality.json").read_text(encoding="utf-8"))
    summary = json.loads((OUTPUT_DIR / "analysis_summary.json").read_text(encoding="utf-8"))
    selected_params = json.loads((OUTPUT_DIR / "selected_parameters.json").read_text(encoding="utf-8"))
    panel = pd.read_csv(OUTPUT_DIR.parent / "data" / "processed" / "quarterly_panel.csv", dtype={"ts_code": str})
    split = pd.read_csv(OUTPUT_DIR / "split_summary.csv")
    model_metrics = pd.read_csv(OUTPUT_DIR / "model_metrics.csv").set_index("model")
    quarterly_ic = pd.read_csv(OUTPUT_DIR / "quarterly_ic.csv")
    backtest_metrics = pd.read_csv(OUTPUT_DIR / "backtest_metrics.csv").set_index("model")
    quarterly_returns = pd.read_csv(OUTPUT_DIR / "quarterly_returns.csv")
    cost_sensitivity = pd.read_csv(OUTPUT_DIR / "cost_sensitivity.csv")
    importance = pd.read_csv(OUTPUT_DIR / "feature_importance.csv")
    best_model = summary["best_model_by_test_mean_rank_ic"]
    best_name = summary["best_model_zh"]
    best_metric = backtest_metrics.loc[best_model]
    best_quarterly = quarterly_returns.loc[quarterly_returns["model"] == best_model].sort_values("quarter_index")

    correlation = panel[[f"x_{feature}" for feature in FEATURE_COLUMNS]].corr(method="spearman")
    mask = np.triu(np.ones(correlation.shape, dtype=bool), k=1)
    pairs = correlation.where(mask).stack().sort_values(key=lambda series: series.abs(), ascending=False)
    top_pair = pairs.index[0]
    top_pair_value = float(pairs.iloc[0])

    doc = Document()
    configure_document(doc)
    doc.core_properties.title = TITLE
    doc.core_properties.subject = "机器学习股票收益排序与季度选股回测课程作业"
    doc.core_properties.author = ""
    doc.core_properties.keywords = "量化交易, 机器学习, 决策树, 随机森林, Rank IC, 回测"

    for _ in range(3):
        spacer = doc.add_paragraph()
        set_paragraph_format(spacer, alignment=WD_ALIGN_PARAGRAPH.CENTER, first_line_chars=None, line_spacing=1.0, space_after=12)
    kicker = doc.add_paragraph()
    set_paragraph_format(kicker, alignment=WD_ALIGN_PARAGRAPH.CENTER, first_line_chars=None, space_after=14)
    set_run_font(kicker.add_run("北京大学工作坊个人作业"), size=11, bold=True, color=GOLD, font=FONT_HEADING)
    title = doc.add_paragraph()
    set_paragraph_format(title, alignment=WD_ALIGN_PARAGRAPH.CENTER, first_line_chars=None, space_after=10)
    set_run_font(title.add_run("TASK6 智能决策者：用机器学习定制专属策略"), size=20, bold=True, color=BLUE, font=FONT_HEADING)
    subtitle = doc.add_paragraph()
    set_paragraph_format(subtitle, alignment=WD_ALIGN_PARAGRAPH.CENTER, first_line_chars=None, space_after=40)
    set_run_font(subtitle.add_run("季度收益排序模型与 Top 30 策略回测"), size=14, bold=True, color=INK, font=FONT_HEADING)
    metadata = doc.add_paragraph()
    set_paragraph_format(metadata, alignment=WD_ALIGN_PARAGRAPH.CENTER, first_line_chars=None, line_spacing=1.5)
    set_run_font(metadata.add_run(f"姓名：____________\n完成日期：2026 年 7 月 18 日\n研究样本：当前中证 300 成分股中等距抽取 {quality['raw_stock_count']} 只"), size=11, color=MUTED)
    notice = doc.add_paragraph()
    set_paragraph_format(notice, alignment=WD_ALIGN_PARAGRAPH.CENTER, first_line_chars=None, space_before=34)
    set_run_font(notice.add_run("本报告仅用于课程研究，不构成投资建议。"), size=9.5, color=MUTED, italic=True)
    doc.add_page_break()

    add_heading(doc, "摘要", 1)
    add_body(
        doc,
        f"本文构建了一套面向股票横截面收益排序的机器学习季度选股流程。研究从抓取日的当前中证 300 成分列表中按代码排序等距抽取 {quality['raw_stock_count']} 只股票，获取 2015 年以来后复权日行情，衍生 {len(FEATURE_COLUMNS)} 个动量、趋势、风险和流动性因子，以未来一个季度收益的横截面百分位为应变量。数据严格按时间划分，先在训练期拟合，在验证期选择超参数，再对最后 {int(best_metric['quarters'])} 个季度采用扩展窗口逐季重训。模型包括岭回归、决策树和随机森林，每季度按预测分数选择前 {TOP_N} 只股票等权持有，并与同一可投资样本的等权平均收益比较。",
    )
    add_body(
        doc,
        f"样本外结果显示，按平均 Rank IC 选择的最佳模型为{best_name}，测试期平均 Rank IC 为 {summary['best_mean_rank_ic']:.3f}，Rank IC 为正的季度比例为 {summary['best_positive_ic_ratio']:.1%}。在单边综合成本 {ONE_WAY_COST:.1%} 的假设下，其 Top {TOP_N} 组合累计净收益为 {summary['best_strategy_cumulative_net_return']:.2%}，同期样本等权平均累计收益为 {summary['benchmark_cumulative_return']:.2%}，最大回撤为 {summary['best_strategy_max_drawdown']:.2%}。这些结果说明模型在本样本中具有一定排序信息，但当前成分股回溯会产生幸存者偏差，且历史 ST、停牌、涨跌停和冲击成本未被完整建模，因此结论只能视为教学型样本外证据。",
    )
    add_callout(doc, "关键词", "机器学习；股票收益排序；决策树；随机森林；Rank IC；Top 30；季度回测")

    add_heading(doc, "目录", 1)
    for item in (
        "1 机器学习交易策略的核心理念",
        "2 常见自变量因子与应变量",
        "3 数据、时点和样本划分",
        "4 模型构建与评价方法",
        "5 交易策略与回测结果",
        "6 模型对比、稳健性与局限",
        "7 结论",
        "参考文献",
        "附录：复现说明与核心代码",
    ):
        add_bullet(doc, item)

    add_heading(doc, "1 机器学习交易策略的核心理念", 1)
    add_heading(doc, "1.1 核心理念", 2)
    add_body(
        doc,
        "机器学习交易策略把每个决策时点可观察到的价格、成交、估值和财务信息转化为特征向量，再学习这些特征与未来收益、未来超额收益或未来收益排序之间的统计映射。模型输出并不直接等于交易，而是要经过股票池筛选、排序、持仓数量、权重、调仓频率、成本和风险约束，才形成可回测的投资组合。本研究的预测目标是下一季度收益排序，因此评价重点是 Rank IC 和 Top 30 组合，而不是只追求收益率点预测误差最小。",
    )
    add_body(
        doc,
        "这一方法的逻辑可概括为：在季度末使用当时已经形成的历史行情构造因子，模型对所有可投资股票给出预测分数，按分数从高到低选择 30 只股票，在下一季度持有并记录收益；随后将新完成的季度加入训练数据，再进行下一次预测。时间顺序是策略可信度的核心。如果把未来价格、未来公告或测试期信息用于因子处理和调参，即使回测曲线很好，也不能视为有效证据。",
    )
    add_heading(doc, "1.2 优点", 2)
    for text in (
        "能够同时处理大量变量，并从非线性关系和因子交互中提取信号；",
        "模型分数天然适合横截面排序，可以直接连接选股和组合构建；",
        "通过扩展窗口或滚动窗口重训，可以对市场结构变化作出一定适应；",
        "决策树规则直观，随机森林通过多棵树平均降低单树方差；",
        "数据、模型、持仓和绩效可以由 Python 一体化复现，便于审计。",
    ):
        add_bullet(doc, text)
    add_heading(doc, "1.3 缺点", 2)
    for text in (
        "金融样本噪声大、非平稳，历史相关关系可能快速衰减；",
        "模型容易过拟合，超参数、因子和时间区间的反复尝试会形成数据挖掘偏差；",
        "未来函数、公告日期错配、幸存者偏差和成分股回溯会显著夸大结果；",
        "复杂模型的可解释性弱，特征重要度不等于因果贡献；",
        "手续费、印花税、滑点、涨跌停和冲击成本会侵蚀纸面收益；",
        "预测指标较好不必然转化为投资收益，组合集中度和换手率同样关键。",
    ):
        add_bullet(doc, text)

    add_heading(doc, "2 常见自变量因子与应变量", 1)
    add_heading(doc, "2.1 常见自变量", 2)
    add_body(
        doc,
        "量化交易中的自变量通常分为动量、估值、质量、成长、规模、风险、流动性和市场状态八类。动量衡量价格趋势；估值描述价格相对盈利或账面价值是否昂贵；质量和成长描述企业盈利能力与扩张速度；规模反映市值风格；风险因子衡量波动、Beta 和回撤；流动性因子衡量换手、成交额和价格冲击；市场状态则描述指数趋势、波动环境和宏观条件。实际模型只能使用决策时点已经公开的数据。",
    )
    factor_rows = [[FEATURE_LABELS[f][0], FEATURE_LABELS[f][1], FEATURE_LABELS[f][2]] for f in FEATURE_COLUMNS]
    add_table_caption(doc, "表 1 本研究自变量因子定义")
    add_table(doc, ["因子组", "变量", "定义"], factor_rows, [1500, 2300, 5225])
    add_heading(doc, "2.2 常见应变量", 2)
    add_body(
        doc,
        "常见应变量包括未来绝对收益、相对基准的未来超额收益、横截面收益排名、未来上涨或跑赢基准的概率、未来波动率、未来最大回撤以及违约或风险事件。回归模型适合预测连续收益或排名，分类模型适合预测方向或是否跑赢。本研究把未来一季度收益先计算为连续变量，再在同一季度内转化为百分位排名 target_rank。这样既保留相对强弱信息，又减少极端收益对模型损失函数的影响。",
    )
    add_callout(doc, "主应变量", "forward_return = 下一季度第二个季度首个交易日开盘价 / 下一季度首个交易日开盘价 - 1；target_rank 为其当季横截面百分位。")

    add_heading(doc, "3 数据、时点和样本划分", 1)
    add_heading(doc, "3.1 数据来源与研究股票池", 2)
    add_body(
        doc,
        f"本研究使用 AkShare 公开接口获取中证指数网站返回的当前中证 300 成分列表，并按股票代码排序等距抽取 {quality['raw_stock_count']} 只作为研究样本，再获取东方财富或腾讯的后复权日行情。原始数据覆盖 {quality['raw_date_min']} 至 {quality['raw_date_max']}，共有 {quality['raw_daily_rows']:,} 条日频记录。后复权价格用于计算跨分红送转的历史收益，并避免长期前复权价格过小引起的四舍五入伪影；当东方财富接口限流时，脚本使用腾讯接口作为备选，并将缺失的换手率或成交额因子填为当季横截面中性值。",
    )
    add_callout(doc, "数据边界", "历史股票池使用当前成分股回溯，不是逐季度真实成分股，存在幸存者和样本选择偏差；报告中的“市场平均”特指当季可用研究样本的等权平均，不等同于全部 A 股或指数官方收益。")
    add_figure(
        doc,
        "figure_1_sample_coverage.png",
        "图 1 每季度模型样本覆盖",
        f"图 1 显示季度面板共覆盖 {quality['panel_quarters']} 个季度，每季度可用股票数介于 {quality['min_quarter_stock_count']} 和 {quality['max_quarter_stock_count']} 之间。随着新上市股票累积满 252 个交易日，样本数可能增加；验证期和测试期的着色区间明确了参数选择与最终评价边界。",
    )
    add_heading(doc, "3.2 决策时点与防泄漏设计", 2)
    add_body(
        doc,
        "每个自然季度最后一个交易日收盘后形成信号，因子只使用截至该日的历史行情。组合在下一季度第一个可用交易日开盘建仓，在再下一个季度第一个可用交易日开盘退出。训练、验证和测试按季度先后排列，不做随机划分；测试期每个季度都只使用该季度之前已经完成的样本重新训练。横截面缩尾、缺失填补和百分位转换只使用同一信号日的股票，不使用未来季度分布。",
    )
    split_rows = [
        [
            row["split"],
            row["quarter_start"],
            row["quarter_end"],
            str(int(row["quarters"])),
            f"{int(row['rows']):,}",
            str(int(row["stocks"])),
        ]
        for _, row in split.iterrows()
    ]
    add_table_caption(doc, "表 2 训练集、验证集和测试集时间划分")
    add_table(doc, ["数据集", "起始季度", "结束季度", "季度数", "样本行", "股票数"], split_rows, [1200, 1500, 1500, 1200, 1700, 1925])
    add_heading(doc, "3.3 数据质量处理", 2)
    add_body(
        doc,
        f"季度面板以“股票代码 + 信号季度”为唯一键，最终包含 {quality['panel_rows']:,} 条记录、{quality['panel_stock_count']} 只股票和 {len(FEATURE_COLUMNS)} 个因子。模型样本要求至少 252 个历史交易日、至少 14 个非缺失因子、正的进出场价格以及合理范围内的未来收益。每个季度内对因子按 1% 和 99% 分位缩尾，以中位数填补缺失，再转换为百分位排名。自动校验确认面板重复主键为 {quality['panel_duplicate_keys']}。",
    )
    add_figure(
        doc,
        "figure_2_factor_correlation.png",
        "图 2 模型因子 Spearman 相关性",
        f"图 2 表明部分动量和趋势变量之间存在较强相关。绝对相关度最高的一组是 {top_pair[0].replace('x_', '')} 与 {top_pair[1].replace('x_', '')}，Spearman 相关系数为 {top_pair_value:.3f}。树模型可在分裂中选择替代变量，岭回归则通过正则化缓解共线性，但重要度仍可能在相关特征之间分散。",
    )

    add_heading(doc, "4 模型构建与评价方法", 1)
    add_heading(doc, "4.1 模型设置", 2)
    add_body(
        doc,
        "岭回归作为线性基准，先对因子标准化，再通过 L2 正则化压缩系数；决策树通过递归分裂拟合非线性规则，并用树深和叶节点最小样本数限制过拟合；随机森林对多个 bootstrap 样本训练决策树，并在每次分裂时抽取部分特征，最后平均预测。三类模型都输出连续分数，以便按股票横截面排序。",
    )
    param_rows = []
    for key in MODEL_ORDER:
        readable = "；".join(f"{name}={value}" for name, value in selected_params[key].items())
        if key == "random_forest":
            readable = "n_estimators=300；" + readable
        param_rows.append([MODEL_NAMES_ZH[key], readable, "验证期平均 Rank IC 最大"])
    add_table_caption(doc, "表 3 验证期选择的模型参数")
    add_table(doc, ["模型", "最终参数", "选择规则"], param_rows, [1600, 4925, 2500])
    add_heading(doc, "4.2 评价指标", 2)
    add_body(
        doc,
        "预测层同时报告 MAE、RMSE、R² 和季度 Rank IC。MAE 与 RMSE 衡量预测排名与真实排名的误差，R² 衡量相对均值预测的解释程度；Rank IC 是预测分数排序与实际未来收益排序的 Spearman 相关系数，更直接对应选股任务。由于金融信号弱、噪声高，R² 可能接近零或为负，但正且相对稳定的 Rank IC 仍可能具有组合意义。策略层再报告收益、波动、Sharpe、最大回撤、Calmar、换手率、跟踪误差和信息比率。",
    )
    model_rows = []
    for key in MODEL_ORDER:
        row = model_metrics.loc[key]
        model_rows.append(
            [
                MODEL_NAMES_ZH[key],
                num(row["mae_rank"]),
                num(row["rmse_rank"]),
                num(row["r2_rank"]),
                num(row["mean_rank_ic"]),
                num(row["icir_quarterly"]),
                pct(row["positive_ic_ratio"], 1),
            ]
        )
    add_table_caption(doc, "表 4 三种模型测试期预测指标")
    add_table(doc, ["模型", "MAE", "RMSE", "R²", "平均 Rank IC", "季度 ICIR", "IC>0 比例"], model_rows, [1300, 1050, 1050, 1050, 1800, 1400, 1375])
    add_figure(
        doc,
        "figure_3_quarterly_rank_ic.png",
        "图 3 三种模型测试期季度 Rank IC",
        f"图 3 展示 {int(best_metric['quarters'])} 个测试季度的排序相关。{best_name}的平均 Rank IC 最高，为 {summary['best_mean_rank_ic']:.3f}，但季度之间仍有明显波动，且只有 {summary['best_positive_ic_ratio']:.1%} 的季度为正。这说明模型可能捕捉到弱信号，却不能把单个季度的结果视为稳定规律。",
    )
    model_excess = quarterly_returns.groupby("model")["net_excess_return"].mean()
    add_figure(
        doc,
        "figure_4_model_excess_return.png",
        "图 4 三种模型 Top 30 组合平均季度净超额收益",
        "图 4 把预测结果进一步转化为投资结果。"
        + "；".join(f"{MODEL_NAMES_ZH[key]}为 {model_excess.loc[key]:+.2%}" for key in MODEL_ORDER)
        + "。预测 Rank IC 与组合超额收益方向可能不同，因为 Top 30 只关注分布最上端，且个股极端收益、集中度与换手成本会放大差异。",
    )
    top_features = importance.head(3)
    add_figure(
        doc,
        "figure_8_random_forest_feature_importance.png",
        "图 8 随机森林前 12 项特征重要度",
        "随机森林平均重要度最高的三项是 "
        + "、".join(f"{row['feature'].replace('x_', '')}（{row['mean_importance']:.3f}）" for _, row in top_features.iterrows())
        + "。该指标反映特征在树分裂中降低不纯度的相对贡献，但在相关因子之间会分散，且不能解释为因果效应。",
    )

    add_heading(doc, "5 交易策略与回测结果", 1)
    add_heading(doc, "5.1 组合规则和市场基准", 2)
    add_body(
        doc,
        f"每个测试季度对可投资股票按模型预测分数降序排列，选择前 {TOP_N} 只股票等权配置，持有到下一次季度调仓。若某季度可选股票少于 {TOP_N}，则对实际可选股票等权。市场基准定义为同一季度所有可投资研究样本股票的等权平均收益，因此策略和基准使用完全相同的进出场价格与股票池。主结果按换手率扣除单边 {ONE_WAY_COST:.1%} 综合成本，并对 0.1%、0.2% 和 0.3% 做敏感性分析。",
    )
    backtest_rows = []
    for key in MODEL_ORDER:
        row = backtest_metrics.loc[key]
        backtest_rows.append(
            [
                MODEL_NAMES_ZH[key],
                pct(row["cumulative_return"]),
                pct(row["annualized_return"]),
                pct(row["annualized_volatility"]),
                num(row["sharpe_rf0"]),
                pct(row["max_drawdown"]),
                pct(row["annualized_excess_return"]),
                num(row["information_ratio"]),
            ]
        )
    add_table_caption(doc, "表 5 三种模型 Top 30 组合回测核心指标（净收益）")
    add_table(doc, ["模型", "累计收益", "年化收益", "年化波动", "Sharpe", "最大回撤", "年化超额", "信息比率"], backtest_rows, [1150, 1150, 1150, 1150, 1000, 1150, 1150, 1125])
    add_figure(
        doc,
        "figure_5_cumulative_nav.png",
        "图 5 最优排序模型策略与市场平均组合累计净值",
        f"测试期内，{best_name} Top {TOP_N} 组合扣费后累计收益为 {best_metric['cumulative_return']:.2%}，样本等权平均组合累计收益为 {best_metric['benchmark_cumulative_return']:.2%}。累计净值展示的是路径和复利结果，不等于统计显著性；较短的 {int(best_metric['quarters'])} 个季度样本容易受到个别行情阶段影响。",
    )
    add_heading(doc, "5.2 逐季度收益", 2)
    quarterly_rows = [
        [
            row["quarter"],
            pct(row["net_return"], signed=True),
            pct(row["benchmark_return"], signed=True),
            pct(row["net_excess_return"], signed=True),
            pct(row["turnover"], 1),
        ]
        for _, row in best_quarterly.iterrows()
    ]
    add_table_caption(doc, f"表 6 {best_name} Top {TOP_N} 组合测试期季度收益")
    add_table(doc, ["信号季度", "策略净收益", "市场平均", "净超额", "换手率"], quarterly_rows, [1700, 1800, 1800, 1800, 1925])
    best_period = best_quarterly.loc[best_quarterly["net_return"].idxmax()]
    worst_period = best_quarterly.loc[best_quarterly["net_return"].idxmin()]
    add_figure(
        doc,
        "figure_6_quarterly_returns.png",
        "图 6 最优模型策略与市场平均组合季度收益",
        f"{best_name}策略表现最好的信号季度为 {best_period['quarter']}，下一持有期净收益 {best_period['net_return']:+.2%}；最差季度为 {worst_period['quarter']}，净收益 {worst_period['net_return']:+.2%}。柱状图显示策略并非每季度都跑赢市场，超额收益具有明显时变性。",
    )
    add_heading(doc, "5.3 风险与交易成本", 2)
    add_figure(
        doc,
        "figure_7_drawdown.png",
        "图 7 最优模型策略与市场平均组合回撤",
        f"{best_name} Top {TOP_N} 组合的最大回撤为 {best_metric['max_drawdown']:.2%}，年化波动率为 {best_metric['annualized_volatility']:.2%}。回撤反映从历史峰值到随后低点的损失，能够揭示累计收益终值掩盖的路径风险。",
    )
    sensitivity_rows = []
    for _, row in cost_sensitivity.iterrows():
        sensitivity_rows.append(
            [
                row["model_zh"],
                pct(row["one_way_cost"], 1),
                pct(row["cumulative_return"]),
                pct(row["annualized_return"]),
                pct(row["annualized_excess_return"]),
                num(row["information_ratio"]),
            ]
        )
    add_table_caption(doc, "表 7 交易成本敏感性分析")
    add_table(doc, ["模型", "单边成本", "累计净收益", "年化净收益", "年化净超额", "信息比率"], sensitivity_rows, [1300, 1400, 1600, 1600, 1600, 1525])
    add_body(
        doc,
        "成本敏感性表说明，高换手策略的结果会随成本假设明显变化。本研究采用统一综合费率便于比较，但并未逐笔区分佣金、印花税和买卖价差，也未估计交易量相对市场深度的冲击成本，因此实盘可实现收益通常低于回测结果。",
    )

    add_heading(doc, "6 模型对比、稳健性与局限", 1)
    add_heading(doc, "6.1 决策树与随机森林比较", 2)
    tree_pred = model_metrics.loc["decision_tree"]
    forest_pred = model_metrics.loc["random_forest"]
    tree_bt = backtest_metrics.loc["decision_tree"]
    forest_bt = backtest_metrics.loc["random_forest"]
    add_body(
        doc,
        f"单棵决策树测试期平均 Rank IC 为 {tree_pred['mean_rank_ic']:.3f}，随机森林为 {forest_pred['mean_rank_ic']:.3f}；相应 Top {TOP_N} 组合年化净收益分别为 {tree_bt['annualized_return']:.2%} 和 {forest_bt['annualized_return']:.2%}。决策树的优势是规则可读、计算快，但它对训练样本扰动敏感，深度过大时容易过拟合。随机森林通过 bootstrap 和特征抽样平均多棵树，一般能降低方差，但模型更难解释、训练成本更高，重要度也可能偏向可分裂点较多或相关的变量。",
    )
    add_heading(doc, "6.2 稳健性判断", 2)
    add_body(
        doc,
        "本研究采用了四项基本稳健性控制：第一，训练、验证和测试严格按时间先后划分；第二，测试期使用扩展窗口逐季重训，而不是一次性用全部数据；第三，同时检查预测 Rank IC 和投资组合收益，避免只用一个指标选模型；第四，对交易成本做多情景分析。尽管如此，测试期只有 8 个季度，无法据此证明长期稳定性；更可靠的研究还应进行历史成分股回溯、行业和风格中性化、不同持仓数量、滚动窗口、不同调仓价格和多时间段检验。",
    )
    add_heading(doc, "6.3 主要局限", 2)
    for caveat in (
        "股票池以当前中证 300 成分股回溯历史，遗漏过去被剔除或退市股票，存在幸存者偏差；",
        "未获得逐季度点时成分、历史 ST 和完整停牌/涨跌停状态，成交可达性被简化；",
        "主模型只使用行情、趋势、风险和流动性因子，没有加入按公告日对齐的财务质量、估值和成长因子；",
        "东方财富限流时使用腾讯数据补充；腾讯记录缺少成交额和换手率，因此相关因子按当季横截面中位数中性填充；",
        "综合成本为情景假设，未精确模拟佣金档位、印花税历史变化、滑点和冲击成本；",
        "测试季度较少，结果受市场阶段和极端个股收益影响，不能解释为因果或未来保证。",
    ):
        add_bullet(doc, caveat)
    add_callout(doc, "总体可信度", "可作为课程中的完整、可复现机器学习排序与回测演示；因股票池和交易可达性限制，真实投资结论应标记为“可分享但必须附带局限”。")

    add_heading(doc, "7 结论", 1)
    add_body(
        doc,
        f"本文完成了从公开行情加载、季度因子衍生、未来收益排序标签设计、时间序列划分、岭回归/决策树/随机森林训练，到 Top {TOP_N} 组合回测和市场平均比较的完整流程。结果显示，最佳测试期平均 Rank IC 模型为{best_name}，说明不同模型对弱金融信号的稳定性存在差异；与此同时，模型预测指标、净超额收益和回撤并不完全同步，策略评价必须把统计排序能力、交易成本和风险路径结合起来。",
    )
    add_body(
        doc,
        "机器学习的价值不在于自动创造确定收益，而在于以统一规则处理多因子信息、形成可检验的排序，并通过严格样本外回测筛除无效假设。下一步若要提升研究可信度，应优先使用逐季度历史股票池，按公告日期加入财务和估值因子，模拟停牌、涨跌停和实际成本，再扩大测试期并进行行业/风格中性化，而不是继续在同一测试集反复调参。",
    )

    add_heading(doc, "参考文献", 1)
    references = (
        "Breiman, L. (2001). Random Forests. Machine Learning, 45, 5-32.",
        "Fama, E. F., & French, K. R. (1993). Common risk factors in the returns on stocks and bonds. Journal of Financial Economics, 33(1), 3-56.",
        "Gu, S., Kelly, B., & Xiu, D. (2020). Empirical Asset Pricing via Machine Learning. Review of Financial Studies, 33(5), 2223-2273.",
        "Hastie, T., Tibshirani, R., & Friedman, J. (2009). The Elements of Statistical Learning (2nd ed.). Springer.",
        "López de Prado, M. (2018). Advances in Financial Machine Learning. Wiley.",
        "AkShare 项目文档与公开接口：中证指数成分列表、A 股历史行情。数据抓取日期见 source_metadata.json。",
    )
    for item in references:
        add_bullet(doc, item)

    add_heading(doc, "附录：复现说明与核心代码", 1)
    add_body(doc, "在仓库根目录依次运行以下命令。原始公开数据已缓存时，可以跳过第一条联网命令。")
    code_lines = [
        r"# 公开数据抓取（需联网且安装 AkShare）",
        r"python .\TASK6\scripts\fetch_data.py --max-stocks 120 --workers 2",
        r"# 因子、模型、回测和图表",
        r".\.venv\Scripts\python.exe .\TASK6\scripts\analyze_strategy.py",
        r"# 生成 DOCX 报告",
        r".\.venv\Scripts\python.exe .\TASK6\scripts\build_report.py",
        r"# 自动数值与交付校验",
        r".\.venv\Scripts\python.exe .\TASK6\scripts\validate_outputs.py",
    ]
    for line in code_lines:
        paragraph = doc.add_paragraph()
        set_paragraph_format(paragraph, alignment=WD_ALIGN_PARAGRAPH.LEFT, line_spacing=1.0, first_line_chars=None)
        shade_paragraph(paragraph, LIGHT_GRAY)
        set_run_font(paragraph.add_run(line), size=8.5, font="Consolas")
    add_body(doc, "完整脚本会保存季度面板、逐股票预测、Top 30 持仓、逐季度收益、模型指标、成本敏感性、图表和校验报告。Notebook 为结果审计与教学伴随文件，不替代脚本。")

    output_path = REPORT_DIR / f"{SUBMISSION_STEM}.docx"
    doc.save(output_path)
    source_notes = [
        "# TASK6 Report Source Notes",
        "",
        "- Primary source: TASK6/data/raw/source_metadata.json",
        "- Reproducible analysis: TASK6/scripts/analyze_strategy.py",
        "- Model evidence: TASK6/outputs/model_metrics.csv and quarterly_ic.csv",
        "- Backtest evidence: TASK6/outputs/quarterly_returns.csv, holdings.csv, backtest_metrics.csv",
        "- Chart map: TASK6/outputs/chart_map.csv",
        "- Named design override: narrative_proposal adapted to A4, SimSun 10.5 pt, 1.5 line spacing, 0 pt paragraph spacing, justified body per assignment instructions.",
        "- Header pattern: restrained editorial cover; running header/footer removed from first page.",
        "- Delivery: native DOCX plus PDF; no website report was requested.",
    ]
    (OUTPUT_DIR / "report_source_notes.md").write_text("\n".join(source_notes), encoding="utf-8")
    print(f"DOCX report created: {output_path}")
    return output_path


def main() -> int:
    build_document()
    return 0


if __name__ == "__main__":
    sys.exit(main())
