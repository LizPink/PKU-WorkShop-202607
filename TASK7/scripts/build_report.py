from __future__ import annotations

import argparse
import json
import math
from pathlib import Path

import numpy as np
import pandas as pd
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Twips


ROOT = Path(__file__).resolve().parents[2]
TASK7_DIR = ROOT / "TASK7"
OUTPUT_DIR = TASK7_DIR / "outputs"
FIGURE_DIR = OUTPUT_DIR / "figures"
REPORT_DIR = TASK7_DIR / "report"
TOC_PAGES_PATH = OUTPUT_DIR / "toc_pages.json"

AUTHOR = "李子平"
REPORT_DATE = "2026 年 7 月 25 日"
TITLE = "JoinQuant 策略部署与模拟交易实战"
SUBTITLE = "月度动量、趋势过滤与风险约束策略的设计、调参与独立复核"
SUBMISSION_STEM = f"{AUTHOR}TASK7"

FONT_BODY = "SimSun"
FONT_HEADING = "SimHei"
FONT_CHART = "Microsoft YaHei"
BODY_SIZE = 10.5
TABLE_SIZE = 8.8
CONTENT_WIDTH_DXA = 9025

INK = "25313A"
BLUE = "2F6B8A"
BLUE_DARK = "244C61"
BLUE_LIGHT = "E8F1F5"
GOLD = "C6912B"
ORANGE = "D77A3D"
PINK = "B45B7C"
OLIVE = "738154"
MUTED = "6B747C"
LIGHT_GRAY = "F4F6F7"
GRID = "B7C1C8"
WHITE = "FFFFFF"
PALETTE = {
    "blue": "#2F6B8A",
    "blue_dark": "#244C61",
    "blue_light": "#B9D4DF",
    "gold": "#C6912B",
    "orange": "#D77A3D",
    "pink": "#B45B7C",
    "olive": "#738154",
    "gray": "#7B858D",
    "light_gray": "#DCE2E5",
    "ink": "#25313A",
}


TOC_ENTRIES = [
    (1, "摘要"),
    (1, "技术摘要"),
    (1, "1 JoinQuant 平台认知与操作流程"),
    (2, "1.1 注册认证的合规边界与验收标准"),
    (2, "1.2 平台界面、研究工具与回测功能"),
    (2, "1.3 数据获取、策略 API 与支持资源"),
    (1, "2 策略模板调整与交易规则设计"),
    (2, "2.1 从单标的模板升级为多标的风险约束策略"),
    (2, "2.2 股票池、信号、仓位与交易规则"),
    (2, "2.3 JoinQuant 代码结构与防错处理"),
    (1, "3 回测设计、参数调整与结果评价"),
    (2, "3.1 数据范围、样本切分与撮合假设"),
    (2, "3.2 参数调整显著改善验证期表现"),
    (2, "3.3 参数面并不平滑，仍需警惕过拟合"),
    (1, "4 模拟交易部署与运行推演"),
    (2, "4.1 官方模拟交易部署步骤"),
    (2, "4.2 锁参版本、监控字段与运行检查"),
    (2, "4.3 锁参观察期取得正收益，但观察窗口仍短"),
    (1, "5 风险暴露、稳健性与实盘差异"),
    (2, "5.1 回撤改善不等于风险消失"),
    (2, "5.2 交易成本与高换手削弱策略优势"),
    (2, "5.3 行业与个股集中度仍是主要风险源"),
    (2, "5.4 尾部损失、模型边界与回测偏差"),
    (1, "6 经验、教训与后续改进"),
    (2, "6.1 平台实施过程中的经验"),
    (2, "6.2 从回测到模拟交易的提交前检查"),
    (1, "7 结论"),
    (1, "参考文献"),
    (1, "附录 核心策略代码与参数说明"),
]


def ensure_dirs() -> None:
    FIGURE_DIR.mkdir(parents=True, exist_ok=True)
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)


def read_inputs() -> dict[str, object]:
    comparison = pd.read_csv(OUTPUT_DIR / "strategy_comparison.csv")
    parameter_search = pd.read_csv(OUTPUT_DIR / "parameter_search.csv")
    daily = pd.read_csv(OUTPUT_DIR / "daily_backtest.csv", parse_dates=["trade_date"]).set_index("trade_date")
    baseline_daily = pd.read_csv(
        OUTPUT_DIR / "baseline_daily_backtest.csv",
        parse_dates=["trade_date"],
    ).set_index("trade_date")
    trades = pd.read_csv(OUTPUT_DIR / "trade_log.csv")
    costs = pd.read_csv(OUTPUT_DIR / "cost_sensitivity.csv")
    stock_exposure = pd.read_csv(OUTPUT_DIR / "stock_exposure.csv")
    industry_exposure = pd.read_csv(OUTPUT_DIR / "industry_exposure.csv")
    universe = pd.read_csv(TASK7_DIR / "data" / "strategy_universe.csv")
    selected = json.loads((OUTPUT_DIR / "selected_parameters.json").read_text(encoding="utf-8"))
    quality = json.loads((OUTPUT_DIR / "data_quality.json").read_text(encoding="utf-8"))
    return {
        "comparison": comparison,
        "parameter_search": parameter_search,
        "daily": daily,
        "baseline_daily": baseline_daily,
        "trades": trades,
        "costs": costs,
        "stock_exposure": stock_exposure,
        "industry_exposure": industry_exposure,
        "universe": universe,
        "selected": selected,
        "quality": quality,
    }


def configure_matplotlib() -> None:
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    plt.rcParams.update(
        {
            "font.sans-serif": [FONT_CHART, "SimHei", "Arial Unicode MS"],
            "axes.unicode_minus": False,
            "figure.facecolor": "white",
            "axes.facecolor": "white",
            "axes.edgecolor": PALETTE["ink"],
            "axes.labelcolor": PALETTE["ink"],
            "text.color": PALETTE["ink"],
            "xtick.color": PALETTE["gray"],
            "ytick.color": PALETTE["gray"],
            "axes.titleweight": "bold",
            "axes.titlesize": 14,
            "axes.labelsize": 10,
            "xtick.labelsize": 8.5,
            "ytick.labelsize": 8.5,
            "legend.fontsize": 8.5,
        }
    )


def polish_axis(ax, *, grid_axis: str = "y") -> None:
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.grid(axis=grid_axis, color="#E4E8EA", linewidth=0.7)
    ax.set_axisbelow(True)


def save_figure(fig, filename: str) -> None:
    fig.savefig(FIGURE_DIR / filename, dpi=220, bbox_inches="tight", facecolor="white")
    import matplotlib.pyplot as plt

    plt.close(fig)


def make_workflow_figure() -> None:
    import matplotlib.pyplot as plt
    from matplotlib.patches import FancyArrowPatch, FancyBboxPatch

    fig, ax = plt.subplots(figsize=(9.2, 3.7))
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 4)
    ax.axis("off")
    stages = [
        ("账号与环境", "本人注册认证\n进入策略列表"),
        ("策略开发", "在线编辑器\nAPI与数据检查"),
        ("历史回测", "设定区间/资金\n快速与完整回测"),
        ("参数锁定", "开发期与验证期\n避免使用观察期"),
        ("模拟交易", "创建模拟账户\n监控订单与风险"),
    ]
    colors = [PALETTE["blue_light"], "#D9E4C6", "#F2D7C6", "#E7D5DE", "#E8E1C3"]
    xs = np.linspace(1, 9, len(stages))
    for i, ((heading, detail), x, color) in enumerate(zip(stages, xs, colors)):
        box = FancyBboxPatch(
            (x - 0.75, 1.05),
            1.5,
            1.9,
            boxstyle="round,pad=0.04,rounding_size=0.12",
            linewidth=1.2,
            edgecolor=PALETTE["blue_dark"],
            facecolor=color,
        )
        ax.add_patch(box)
        ax.text(x, 2.35, heading, ha="center", va="center", fontsize=11.5, fontweight="bold")
        ax.text(x, 1.65, detail, ha="center", va="center", fontsize=9.2, linespacing=1.45)
        if i < len(stages) - 1:
            ax.add_patch(
                FancyArrowPatch(
                    (x + 0.78, 2.0),
                    (xs[i + 1] - 0.78, 2.0),
                    arrowstyle="-|>",
                    mutation_scale=14,
                    linewidth=1.4,
                    color=PALETTE["gray"],
                )
            )
    ax.set_title("JoinQuant 策略从开发到模拟交易的完整流程", pad=8)
    ax.text(5, 0.35, "身份认证由本人完成；代码、参数、回测与风险监控应留存可复核记录", ha="center", fontsize=9, color=PALETTE["gray"])
    save_figure(fig, "figure_1_platform_workflow.png")


def rebased_nav(frame: pd.DataFrame, start: str) -> pd.Series:
    sample = frame.loc[frame.index >= pd.Timestamp(start), "strategy_nav"].copy()
    return sample / sample.iloc[0]


def make_nav_figure(data: dict[str, object]) -> None:
    import matplotlib.pyplot as plt

    daily = data["daily"]
    baseline = data["baseline_daily"]
    start = "2024-01-01"
    final_nav = rebased_nav(daily, start)
    baseline_nav = rebased_nav(baseline, start)
    benchmark = daily.loc[daily.index >= pd.Timestamp(start), "benchmark_nav"].copy()
    benchmark = benchmark / benchmark.iloc[0]

    fig, ax = plt.subplots(figsize=(9.2, 4.7))
    ax.plot(final_nav.index, final_nav, color=PALETTE["blue"], linewidth=2.2, label="锁定参数策略")
    ax.plot(baseline_nav.index, baseline_nav, color=PALETTE["orange"], linewidth=1.5, linestyle="--", label="基础模板")
    ax.plot(benchmark.index, benchmark, color=PALETTE["ink"], linewidth=1.4, linestyle=":", label="十股等权基准")
    split = pd.Timestamp("2026-01-01")
    ax.axvline(split, color=PALETTE["gold"], linewidth=1.2)
    ax.text(split, ax.get_ylim()[1], " 锁参观察期", color=PALETTE["gold"], va="top", fontsize=9)
    ax.set_ylabel("累计净值（2024年初=1）")
    ax.set_title("验证期与锁参观察期的累计净值")
    ax.legend(frameon=False, ncol=3, loc="upper left")
    polish_axis(ax)
    save_figure(fig, "figure_2_nav_comparison.png")


def make_period_return_figure(data: dict[str, object]) -> None:
    import matplotlib.pyplot as plt

    comparison = data["comparison"]
    periods = ["验证期", "模拟观察期"]
    categories = ["基础模板", "锁定参数", "十股等权基准"]
    values = {category: [] for category in categories}
    for period in periods:
        base = comparison.loc[
            (comparison["strategy"] == "基础模板") & (comparison["period"] == period)
        ].iloc[0]
        selected = comparison.loc[
            (comparison["strategy"] == "锁定参数") & (comparison["period"] == period)
        ].iloc[0]
        values["基础模板"].append(base["cumulative_return"])
        values["锁定参数"].append(selected["cumulative_return"])
        values["十股等权基准"].append(selected["benchmark_return"])

    x = np.arange(len(periods))
    width = 0.24
    fig, ax = plt.subplots(figsize=(8.8, 4.5))
    colors = [PALETTE["orange"], PALETTE["blue"], PALETTE["ink"]]
    for offset, (category, color) in enumerate(zip(categories, colors)):
        bars = ax.bar(x + (offset - 1) * width, values[category], width, color=color, label=category)
        for bar, value in zip(bars, values[category]):
            ax.text(
                bar.get_x() + bar.get_width() / 2,
                value + (0.018 if value >= 0 else -0.035),
                f"{value:.1%}",
                ha="center",
                va="bottom" if value >= 0 else "top",
                fontsize=9,
            )
    ax.axhline(0, color=PALETTE["ink"], linewidth=0.8)
    ax.set_xticks(x, periods)
    ax.set_ylabel("累计收益率")
    ax.set_title("基础模板、锁定参数与基准的分期收益")
    ax.legend(frameon=False, ncol=3, loc="upper left")
    polish_axis(ax)
    save_figure(fig, "figure_3_period_returns.png")


def make_parameter_heatmap(data: dict[str, object]) -> None:
    import matplotlib.pyplot as plt

    search = data["parameter_search"]
    pivot = search.pivot_table(
        index="momentum_window",
        columns="trend_window",
        values="validation_sharpe_ratio",
        aggfunc="mean",
    ).sort_index().sort_index(axis=1)
    fig, ax = plt.subplots(figsize=(7.8, 4.8))
    image = ax.imshow(pivot.values, cmap="Blues", aspect="auto")
    for i in range(pivot.shape[0]):
        for j in range(pivot.shape[1]):
            value = pivot.iloc[i, j]
            ax.text(j, i, f"{value:.2f}", ha="center", va="center", color="white" if value > pivot.values.mean() else PALETTE["ink"], fontsize=10)
    ax.set_xticks(range(len(pivot.columns)), [f"{int(x)}日" for x in pivot.columns])
    ax.set_yticks(range(len(pivot.index)), [f"{int(x)}日" for x in pivot.index])
    ax.set_xlabel("趋势均线窗口")
    ax.set_ylabel("动量窗口")
    ax.set_title("参数网格的验证期平均夏普比率")
    cbar = fig.colorbar(image, ax=ax, shrink=0.82)
    cbar.set_label("夏普比率")
    save_figure(fig, "figure_4_parameter_heatmap.png")


def drawdown(series: pd.Series) -> pd.Series:
    nav = (1.0 + series.fillna(0.0)).cumprod()
    return nav / nav.cummax() - 1.0


def make_drawdown_figure(data: dict[str, object]) -> None:
    import matplotlib.pyplot as plt

    daily = data["daily"].loc["2024-01-01":]
    baseline = data["baseline_daily"].loc["2024-01-01":]
    selected_dd = drawdown(daily["strategy_return"])
    baseline_dd = drawdown(baseline["strategy_return"])
    benchmark_dd = drawdown(daily["benchmark_return"])
    fig, ax = plt.subplots(figsize=(9.2, 4.5))
    ax.plot(selected_dd.index, selected_dd, color=PALETTE["blue"], linewidth=2.0, label="锁定参数策略")
    ax.plot(baseline_dd.index, baseline_dd, color=PALETTE["orange"], linewidth=1.4, linestyle="--", label="基础模板")
    ax.plot(benchmark_dd.index, benchmark_dd, color=PALETTE["ink"], linewidth=1.2, linestyle=":", label="十股等权基准")
    ax.fill_between(selected_dd.index, selected_dd.values, 0, color=PALETTE["blue"], alpha=0.12)
    ax.set_ylabel("回撤")
    ax.set_title("验证期与锁参观察期的回撤路径")
    ax.legend(frameon=False, ncol=3, loc="lower left")
    polish_axis(ax)
    save_figure(fig, "figure_5_drawdown.png")


def make_cost_figure(data: dict[str, object]) -> None:
    import matplotlib.pyplot as plt

    costs = data["costs"].copy()
    fig, ax = plt.subplots(figsize=(7.8, 4.4))
    bars = ax.bar(costs["scenario"], costs["cumulative_return"], color=[PALETTE["blue_light"], PALETTE["blue"], PALETTE["blue_dark"]])
    for bar, value in zip(bars, costs["cumulative_return"]):
        ax.text(bar.get_x() + bar.get_width() / 2, value + 0.012, f"{value:.1%}", ha="center", fontsize=10)
    ax.set_ylim(0, max(costs["cumulative_return"]) * 1.18)
    ax.set_ylabel("累计收益率")
    ax.set_title("验证期至锁参观察期的交易成本敏感性")
    polish_axis(ax)
    save_figure(fig, "figure_6_cost_sensitivity.png")


def make_exposure_figure(data: dict[str, object]) -> None:
    import matplotlib.pyplot as plt

    industry = data["industry_exposure"].sort_values("average_weight", ascending=True)
    y = np.arange(len(industry))
    fig, ax = plt.subplots(figsize=(8.5, 4.8))
    ax.barh(y - 0.17, industry["average_weight"], height=0.32, color=PALETTE["blue"], label="全观察期平均权重")
    ax.barh(y + 0.17, industry["max_weight"], height=0.32, color=PALETTE["gold"], label="单日最高权重")
    ax.set_yticks(y, industry["industry"])
    ax.set_xlabel("组合权重")
    ax.set_title("锁参观察期的行业风险暴露")
    ax.legend(frameon=False, loc="lower right")
    polish_axis(ax, grid_axis="x")
    save_figure(fig, "figure_7_industry_exposure.png")


def make_tail_figure(data: dict[str, object]) -> None:
    import matplotlib.pyplot as plt

    returns = data["daily"].loc["2026-01-01":"2026-07-10", "strategy_return"].dropna()
    var_95 = returns.quantile(0.05)
    cvar_95 = returns[returns <= var_95].mean()
    fig, ax = plt.subplots(figsize=(8.5, 4.5))
    ax.hist(returns, bins=18, color=PALETTE["blue_light"], edgecolor=PALETTE["blue_dark"], linewidth=0.8)
    ax.axvline(var_95, color=PALETTE["orange"], linewidth=1.8, label=f"95% VaR：{var_95:.2%}")
    ax.axvline(cvar_95, color=PALETTE["pink"], linewidth=1.8, linestyle="--", label=f"95% CVaR：{cvar_95:.2%}")
    ax.set_xlabel("日收益率")
    ax.set_ylabel("交易日数量")
    ax.set_title("锁参观察期的日收益分布与左尾风险")
    ax.legend(frameon=False)
    polish_axis(ax)
    save_figure(fig, "figure_8_tail_risk.png")


def make_figures(data: dict[str, object]) -> None:
    configure_matplotlib()
    make_workflow_figure()
    make_nav_figure(data)
    make_period_return_figure(data)
    make_parameter_heatmap(data)
    make_drawdown_figure(data)
    make_cost_figure(data)
    make_exposure_figure(data)
    make_tail_figure(data)


def set_run_font(run, *, size: float = BODY_SIZE, bold: bool | None = None, color: str = INK, font: str = FONT_BODY, italic: bool | None = None) -> None:
    run.font.name = font
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:ascii"), font)
    rpr.rFonts.set(qn("w:hAnsi"), font)
    rpr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_paragraph_format(paragraph, *, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY, line_spacing: float = 1.5, keep_with_next: bool = False, first_line_chars: float | None = 2, space_before: float = 0, space_after: float = 0, page_break_before: bool = False) -> None:
    fmt = paragraph.paragraph_format
    fmt.alignment = alignment
    fmt.space_before = Pt(space_before)
    fmt.space_after = Pt(space_after)
    fmt.line_spacing = line_spacing
    fmt.keep_with_next = keep_with_next
    fmt.widow_control = True
    fmt.page_break_before = page_break_before
    if first_line_chars is not None:
        fmt.first_line_indent = Pt(BODY_SIZE * first_line_chars)


def add_body(doc: Document, text: str, *, bold_lead: str | None = None) -> object:
    paragraph = doc.add_paragraph()
    set_paragraph_format(paragraph)
    if bold_lead and text.startswith(bold_lead):
        set_run_font(paragraph.add_run(bold_lead), bold=True)
        set_run_font(paragraph.add_run(text[len(bold_lead) :]))
    else:
        set_run_font(paragraph.add_run(text))
    return paragraph


def add_heading(doc: Document, text: str, level: int = 1, *, page_break_before: bool = False, compact: bool = False) -> object:
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    if compact:
        space_before = 3 if level == 1 else 2
        space_after = 2 if level == 1 else 1
    else:
        space_before = 10 if level == 1 else 6
        space_after = 4
    set_paragraph_format(
        paragraph,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
        keep_with_next=True,
        first_line_chars=None,
        space_before=space_before,
        space_after=space_after,
        page_break_before=page_break_before,
    )
    set_run_font(
        paragraph.add_run(text),
        size=15 if level == 1 else 12 if level == 2 else 10.5,
        bold=True,
        color=BLUE if level == 1 else INK,
        font=FONT_HEADING,
    )
    return paragraph


def add_bullet(doc: Document, text: str) -> object:
    paragraph = doc.add_paragraph(style="List Bullet")
    set_paragraph_format(paragraph, first_line_chars=None)
    paragraph.paragraph_format.left_indent = Cm(0.74)
    paragraph.paragraph_format.first_line_indent = Cm(-0.37)
    set_run_font(paragraph.add_run(text))
    return paragraph


def shade_paragraph(paragraph, fill: str) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    shd = ppr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        ppr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_callout(doc: Document, label: str, text: str, *, fill: str = BLUE_LIGHT) -> object:
    paragraph = doc.add_paragraph()
    set_paragraph_format(paragraph, alignment=WD_ALIGN_PARAGRAPH.LEFT, first_line_chars=None, space_before=4, space_after=4)
    shade_paragraph(paragraph, fill)
    set_run_font(paragraph.add_run(f"{label}："), bold=True, color=BLUE)
    set_run_font(paragraph.add_run(text))
    return paragraph


def add_code_block(doc: Document, code: str) -> object:
    paragraph = doc.add_paragraph()
    set_paragraph_format(
        paragraph,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
        first_line_chars=None,
        line_spacing=1.0,
        space_before=3,
        space_after=3,
    )
    shade_paragraph(paragraph, LIGHT_GRAY)
    for index, line in enumerate(code.strip().splitlines()):
        if index:
            paragraph.add_run("\n")
        set_run_font(paragraph.add_run(line), size=9.0, color=INK, font=FONT_BODY)
    return paragraph


def ensure_child(parent, tag: str):
    child = parent.find(qn(tag))
    if child is None:
        child = OxmlElement(tag)
        parent.append(child)
    return child


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = ensure_child(tc_pr, "w:tcMar")
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        margin = ensure_child(tc_mar, f"w:{side}")
        margin.set(qn("w:w"), str(value))
        margin.set(qn("w:type"), "dxa")


def set_row_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_table_geometry(table, widths_dxa: list[int], *, indent_dxa: int = 120) -> None:
    if sum(widths_dxa) != CONTENT_WIDTH_DXA:
        raise ValueError(f"Table widths must sum to {CONTENT_WIDTH_DXA}, got {sum(widths_dxa)}")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = ensure_child(tbl_pr, "w:tblW")
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = ensure_child(tbl_pr, "w:tblInd")
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = ensure_child(tbl_pr, "w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        row.height = None
        set_row_cant_split(row)
        for col_index, cell in enumerate(row.cells):
            width = widths_dxa[col_index]
            cell.width = Twips(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = ensure_child(tc_pr, "w:tcW")
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def add_table_caption(doc: Document, text: str) -> object:
    paragraph = doc.add_paragraph(style="Caption")
    set_paragraph_format(paragraph, alignment=WD_ALIGN_PARAGRAPH.CENTER, keep_with_next=True, first_line_chars=None, space_before=4, space_after=4)
    set_run_font(paragraph.add_run(text), size=9.5, bold=True)
    return paragraph


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths_dxa: list[int], *, alignments: list[object] | None = None, font_size: float = TABLE_SIZE) -> object:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    repeat_table_header(table.rows[0])
    alignments = alignments or [WD_ALIGN_PARAGRAPH.LEFT] + [WD_ALIGN_PARAGRAPH.CENTER] * (len(headers) - 1)
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, BLUE_LIGHT)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        set_paragraph_format(paragraph, alignment=WD_ALIGN_PARAGRAPH.CENTER, first_line_chars=None, line_spacing=1.0, keep_with_next=True)
        set_run_font(paragraph.add_run(header), size=font_size, bold=True)
    for row_index, row_data in enumerate(rows):
        row = table.add_row()
        for index, value in enumerate(row_data):
            cell = row.cells[index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            set_paragraph_format(paragraph, alignment=alignments[index], first_line_chars=None, line_spacing=1.0)
            set_run_font(paragraph.add_run(str(value)), size=font_size)
    set_table_geometry(table, widths_dxa)
    spacer = doc.add_paragraph()
    set_paragraph_format(spacer, first_line_chars=None, line_spacing=1.0, space_after=2)
    return table


def add_figure(doc: Document, filename: str, caption: str, interpretation: str, *, width_cm: float = 15.2) -> None:
    image_path = FIGURE_DIR / filename
    paragraph = doc.add_paragraph()
    set_paragraph_format(paragraph, alignment=WD_ALIGN_PARAGRAPH.CENTER, keep_with_next=True, first_line_chars=None, space_before=4)
    inline_shape = paragraph.add_run().add_picture(str(image_path), width=Cm(width_cm))
    inline_shape._inline.docPr.set("descr", caption)
    inline_shape._inline.docPr.set("title", caption)
    caption_paragraph = doc.add_paragraph(style="Caption")
    set_paragraph_format(caption_paragraph, alignment=WD_ALIGN_PARAGRAPH.CENTER, keep_with_next=True, first_line_chars=None, space_after=2)
    set_run_font(caption_paragraph.add_run(caption), size=9.5, bold=True)
    add_body(doc, interpretation)


def add_page_number(paragraph) -> None:
    set_paragraph_format(paragraph, alignment=WD_ALIGN_PARAGRAPH.CENTER, first_line_chars=None, line_spacing=1.0)
    set_run_font(paragraph.add_run("第 "), size=9, color=MUTED)
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_run_font(run, size=9, color=MUTED)
    set_run_font(paragraph.add_run(" 页"), size=9, color=MUTED)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)
    section.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    normal.font.name = FONT_BODY
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT_BODY)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_BODY)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY)
    normal.font.size = Pt(BODY_SIZE)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    heading_tokens = {
        "Heading 1": (15, BLUE, 10, 4),
        "Heading 2": (12, INK, 6, 3),
        "Heading 3": (10.5, BLUE, 4, 2),
    }
    for style_name, (size, color, before, after) in heading_tokens.items():
        style = doc.styles[style_name]
        style.font.name = FONT_HEADING
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT_HEADING)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_HEADING)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_HEADING)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.name = FONT_BODY
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY)
        style.font.size = Pt(BODY_SIZE)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.line_spacing = 1.5

    caption = doc.styles["Caption"]
    caption.font.name = FONT_BODY
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY)
    caption.font.size = Pt(9.5)
    caption.font.color.rgb = RGBColor.from_string(INK)
    caption.paragraph_format.space_before = Pt(0)
    caption.paragraph_format.space_after = Pt(0)

    header = section.header
    hp = header.paragraphs[0]
    set_paragraph_format(hp, alignment=WD_ALIGN_PARAGRAPH.RIGHT, first_line_chars=None, line_spacing=1.0)
    set_run_font(hp.add_run("TASK7 · JoinQuant 策略部署与模拟交易实战"), size=9, color=MUTED)
    add_page_number(section.footer.paragraphs[0])


def add_cover(doc: Document) -> None:
    for _ in range(4):
        paragraph = doc.add_paragraph()
        set_paragraph_format(paragraph, alignment=WD_ALIGN_PARAGRAPH.CENTER, first_line_chars=None, line_spacing=1.0, space_after=12)
    kicker = doc.add_paragraph()
    set_paragraph_format(kicker, alignment=WD_ALIGN_PARAGRAPH.CENTER, first_line_chars=None, space_after=16)
    set_run_font(kicker.add_run("北京大学工作坊个人作业 · TASK7 实战推演"), size=11, bold=True, color=GOLD, font=FONT_HEADING)
    title = doc.add_paragraph()
    set_paragraph_format(title, alignment=WD_ALIGN_PARAGRAPH.CENTER, first_line_chars=None, space_after=8)
    set_run_font(title.add_run(TITLE), size=24, bold=True, color=BLUE, font=FONT_HEADING)
    subtitle = doc.add_paragraph()
    set_paragraph_format(subtitle, alignment=WD_ALIGN_PARAGRAPH.CENTER, first_line_chars=None, space_after=54)
    set_run_font(subtitle.add_run(SUBTITLE), size=14, bold=True, color=INK, font=FONT_HEADING)
    meta = doc.add_paragraph()
    set_paragraph_format(meta, alignment=WD_ALIGN_PARAGRAPH.CENTER, first_line_chars=None, line_spacing=1.8)
    set_run_font(meta.add_run(f"作者：{AUTHOR}\n完成日期：{REPORT_DATE}\n研究平台：JoinQuant（聚宽）"), size=11, color=MUTED)
    note = doc.add_paragraph()
    set_paragraph_format(note, alignment=WD_ALIGN_PARAGRAPH.CENTER, first_line_chars=None, space_before=42)
    set_run_font(note.add_run("本报告仅用于课程学习与研究，不构成投资建议。"), size=9.5, italic=True, color=MUTED)
    doc.add_page_break()


def load_toc_pages() -> dict[str, int | str]:
    if TOC_PAGES_PATH.exists():
        return json.loads(TOC_PAGES_PATH.read_text(encoding="utf-8"))
    return {}


def add_static_toc(doc: Document, pages: dict[str, int | str]) -> None:
    add_heading(doc, "目录", 1)
    table = doc.add_table(rows=0, cols=2)
    table.autofit = False
    for level, title in TOC_ENTRIES:
        row = table.add_row()
        set_row_cant_split(row)
        left, right = row.cells
        left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p_left = left.paragraphs[0]
        set_paragraph_format(p_left, alignment=WD_ALIGN_PARAGRAPH.LEFT, first_line_chars=None, line_spacing=1.2, space_after=1)
        p_left.paragraph_format.left_indent = Cm(0 if level == 1 else 0.65)
        set_run_font(p_left.add_run(title), size=10.0 if level == 1 else 9.5, bold=level == 1)
        p_right = right.paragraphs[0]
        set_paragraph_format(p_right, alignment=WD_ALIGN_PARAGRAPH.RIGHT, first_line_chars=None, line_spacing=1.2, space_after=1)
        set_run_font(p_right.add_run(str(pages.get(title, "—"))), size=10.0 if level == 1 else 9.5, bold=level == 1)
    table.style = "Table Grid"
    set_table_geometry(table, [8150, 875], indent_dxa=0)
    borders = ensure_child(table._tbl.tblPr, "w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = ensure_child(borders, f"w:{edge}")
        node.set(qn("w:val"), "nil")
    doc.add_page_break()


def pct(value: float, digits: int = 1, signed: bool = False) -> str:
    if pd.isna(value):
        return "—"
    prefix = "+" if signed and value > 0 else ""
    return f"{prefix}{value:.{digits}%}"


def num(value: float, digits: int = 2) -> str:
    return "—" if pd.isna(value) else f"{value:.{digits}f}"


def metric_row(comparison: pd.DataFrame, strategy: str, period: str) -> pd.Series:
    return comparison.loc[
        (comparison["strategy"] == strategy) & (comparison["period"] == period)
    ].iloc[0]


def build_document(data: dict[str, object]) -> Path:
    comparison = data["comparison"]
    search = data["parameter_search"]
    trades = data["trades"]
    costs = data["costs"]
    stock_exposure = data["stock_exposure"]
    industry_exposure = data["industry_exposure"]
    universe = data["universe"]
    quality = data["quality"]

    validation = metric_row(comparison, "锁定参数", "验证期")
    paper = metric_row(comparison, "锁定参数", "模拟观察期")
    baseline_validation = metric_row(comparison, "基础模板", "验证期")
    baseline_paper = metric_row(comparison, "基础模板", "模拟观察期")

    doc = Document()
    configure_document(doc)
    doc.core_properties.title = f"{TITLE}：{SUBTITLE}"
    doc.core_properties.subject = "TASK7 实战推演：策略实盘部署与交易实战"
    doc.core_properties.author = AUTHOR
    doc.core_properties.keywords = "JoinQuant, 聚宽, 动量策略, 趋势过滤, 参数优化, 模拟交易, 风险暴露"

    add_cover(doc)
    add_heading(doc, "摘要", 1)
    abstract = (
        "本文围绕 JoinQuant 平台的策略开发、历史回测、参数调整和模拟交易部署，设计一套月度动量排序、长均线过滤、逆波动率配置与跟踪止损相结合的多标的策略。研究使用十只跨行业 A 股的前复权日线数据，将样本划分为开发期、验证期和锁参观察期，并采用佣金、卖出印花税、滑点与整手约束进行独立复核。参数由60日动量、120日趋势和12%止损调整为20日动量、60日趋势和8%止损后，验证期累计收益由-9.10%改善为38.76%，最大回撤由27.58%降至14.92%；2026年锁参观察期收益为9.65%，略高于十股等权基准。结果同时暴露出高换手、行业集中和短期样本不足等问题，说明策略进入模拟交易后仍应以执行稳定性和风险约束为主要检验目标。"
    )
    add_body(doc, abstract)
    keyword = doc.add_paragraph()
    set_paragraph_format(keyword, first_line_chars=None)
    set_run_font(keyword.add_run("关键词："), bold=True)
    set_run_font(keyword.add_run("JoinQuant；动量策略；趋势过滤；参数优化；模拟交易；风险暴露"))

    add_heading(doc, "技术摘要", 1)
    add_bullet(doc, f"锁定参数为20日动量、60日趋势均线、持有前3名和8%跟踪止损，组合目标仓位90%，单股目标权重上限40%。")
    add_bullet(doc, f"验证期累计收益{pct(validation['cumulative_return'])}、夏普比率{num(validation['sharpe_ratio'])}、最大回撤{pct(validation['max_drawdown'])}；收益为正，但仍低于十股等权基准{pct(validation['benchmark_return'])}。")
    add_bullet(doc, f"2026年1月初至7月10日的锁参观察期收益{pct(paper['cumulative_return'])}，比等权基准高{pct(paper['excess_return'])}，但年化换手约{num(paper['annual_turnover'], 1)}倍，短期结果不能外推为长期实盘能力。")
    add_callout(
        doc,
        "证据边界",
        "本报告完成了 JoinQuant 兼容代码、参数方案和独立日频回测。账号注册、实名认证及在个人账户中创建官方回测/模拟交易涉及本人身份和登录状态，必须由作者本人在平台完成；本文没有虚构回测ID、模拟交易ID或长期模拟账户业绩。",
        fill=LIGHT_GRAY,
    )
    add_static_toc(doc, load_toc_pages())

    add_heading(doc, "1 JoinQuant 平台认知与操作流程", 1)
    add_heading(doc, "1.1 注册认证的合规边界与验收标准", 2)
    add_body(
        doc,
        "JoinQuant 官方新手指引说明，用户可从网站右上角进入注册页，按照页面提示建立个人账号；登录后可从“策略”或“编写策略”进入策略列表。实名认证可能涉及手机号、证件与验证码，这些信息不应写入课程报告，也不能交由第三方保存。对本任务而言，账号环节的验收标准不是截图展示敏感信息，而是本人能够登录、进入策略列表、新建策略并看到回测参数区域。",
    )
    add_callout(
        doc,
        "提交前本人确认",
        "确认个人账号已完成平台要求的认证，并能进入策略编辑、完整回测和模拟交易页面。PDF只保留操作说明，不展示证件号码、手机号、验证码、密码或账户令牌。",
    )
    add_heading(doc, "1.2 平台界面、研究工具与回测功能", 2)
    add_body(
        doc,
        "平台的主要工作区可以按“策略管理—在线编辑—回测设置—结果分析—模拟交易”理解。在线编辑器使用 Python 语法，官方指引称策略会自动保存，也可以手动保存；回测前需设置开始日期、结束日期、初始资金与频率。快速回测用于编译和初步排错，完整回测用于查看更完整的收益、风险、持仓和订单结果。研究模块则适合数据探索、参数批量分析和通过 create_backtest/get_backtest 管理回测结果。",
    )
    add_table_caption(doc, "表1 JoinQuant 平台模块及本任务中的用途")
    add_table(
        doc,
        ["模块", "主要功能", "本任务用途"],
        [
            ["策略列表", "新建、复制和管理策略", "保存基础模板与锁参版本"],
            ["在线编辑器", "编写 Python 策略并查看日志", "实现信号、仓位、止损和监控"],
            ["回测设置", "日期、资金、频率和基准", "统一不同参数的比较口径"],
            ["回测结果", "收益、风险、持仓、订单和记录", "核对收益路径与风险暴露"],
            ["投资研究", "Notebook、数据分析和批量回测", "进行参数网格与结果整理"],
            ["模拟交易", "按实时行情运行策略", "观察订单、成交、持仓与异常"],
            ["帮助与社区", "API 文档、教程和问题检索", "核对函数语义与平台限制"],
        ],
        [1300, 3400, 4325],
    )
    add_figure(
        doc,
        "figure_1_platform_workflow.png",
        "图1 JoinQuant 策略从开发到模拟交易的完整流程",
        "图1强调，模拟交易不是历史回测后的装饰性步骤，而是参数冻结后的新阶段。只有在开发期和验证期完成代码检查、成本设置与风险约束后，才应建立模拟交易；进入模拟阶段后，不应根据每天的盈亏频繁修改参数，否则观察结果会失去独立性。",
    )
    add_heading(doc, "1.3 数据获取、策略 API 与支持资源", 2)
    add_body(
        doc,
        "本策略使用 attribute_history 获取单一标的的历史收盘价，并利用 get_current_data 检查停牌、ST、涨跌停和当日可交易状态。initialize 负责设置基准、真实价格模式、交易成本、滑点和定时任务；run_monthly 触发月度调仓，run_daily 触发止损与收盘后监控；order_target_value 将持仓调整至目标市值；record 记录仓位、现金比例与持仓数量。官方 API 文档特别指出，日频 attribute_history 不包含当日数据，这使“用昨日信号、今日开盘执行”的设计能够避免直接使用当日收盘价下单。",
    )
    add_table_caption(doc, "表2 核心 API、用途与主要防错点")
    add_table(
        doc,
        ["API", "用途", "防错点"],
        [
            ["set_benchmark", "设置收益与风险比较基准", "只能在 initialize 中调用"],
            ["set_order_cost / set_slippage", "设置佣金、税费与滑点", "成本必须与本地复核口径一致"],
            ["attribute_history", "获取历史价格和收益序列", "日频不含当天，避免跨日缓存"],
            ["get_current_data", "检查停牌、ST和涨跌停", "返回对象仅在当天有效"],
            ["run_daily / run_monthly", "安排止损、调仓与记录", "函数必须是全局函数"],
            ["order_target_value", "调整至目标持仓市值", "可能因停牌、涨跌停或现金不足失败"],
            ["record", "在结果图中记录自定义风险变量", "日频展示最后一次记录值"],
        ],
        [1850, 3150, 4025],
    )

    add_heading(doc, "2 策略模板调整与交易规则设计", 1)
    add_heading(doc, "2.1 从单标的模板升级为多标的风险约束策略", 2)
    add_body(
        doc,
        "JoinQuant 的双均线示例适合学习 initialize、历史数据和下单函数，但单标的全仓交易会把策略结果高度绑定于个股。本文保留“趋势过滤—条件成立时持有”的模板思想，将其扩展为固定十股股票池的横截面动量选择，并加入逆波动率权重、目标仓位、单股上限、跟踪止损和交易状态检查。调整后的策略不追求每天交易，而是在月初完成组合重构，日内仅执行风险退出。",
    )
    add_heading(doc, "2.2 股票池、信号、仓位与交易规则", 2)
    universe_rows = [
        [row.joinquant_code, row.stock_name, row.industry]
        for row in universe.itertuples(index=False)
    ]
    add_table_caption(doc, "表3 固定研究股票池")
    add_table(doc, ["JoinQuant代码", "证券简称", "行业"], universe_rows, [2600, 2600, 3825])
    add_body(
        doc,
        "月度调仓日先使用前一交易日及更早的数据计算信号。动量定义为最近收盘价相对20个交易日前收盘价的涨幅；趋势过滤要求最近收盘价高于60日均线；合格股票按动量从高到低排序，选择前3名。对入选股票按20日年化波动率的倒数分配权重，组合目标仓位为90%，单股目标权重不超过40%。持仓期间记录自建仓以来的历史高点，若上一日收盘价较高点回撤8%，则在下一交易日开盘尝试清仓。",
    )
    add_table_caption(doc, "表4 基础模板与锁参版本的交易参数")
    add_table(
        doc,
        ["参数", "基础模板", "锁参版本", "设计含义"],
        [
            ["动量窗口", "60日", "20日", "缩短反应时间"],
            ["趋势均线", "120日", "60日", "更快识别中期趋势"],
            ["持仓数量", "3只", "3只", "在分散与信号强度间折中"],
            ["跟踪止损", "12%", "8%", "更早控制单笔回撤"],
            ["波动率窗口", "20日", "20日", "逆波动率定权"],
            ["目标总仓位", "90%", "90%", "保留现金缓冲"],
            ["单股目标上限", "40%", "40%", "约束集中度"],
            ["调仓频率", "月度", "月度", "降低无必要的日频换仓"],
        ],
        [1800, 1500, 1500, 4225],
    )
    add_heading(doc, "2.3 JoinQuant 代码结构与防错处理", 2)
    add_body(
        doc,
        "代码采用四个定时入口：开盘前清空当日状态，开盘检查跟踪止损，每月第一个交易日开盘后5分钟调仓，收盘后记录风险变量。调仓遵循“先卖后买”，避免因旧仓占用资金导致买单失败；对停牌、ST、名称含退市标识和涨停股票进行过滤；订单返回为空时不假设成交成功。参数通过全局变量集中管理，便于在研究模块中使用 extras 批量传入不同配置。",
    )
    add_code_block(
        doc,
        """
set_benchmark('000300.XSHG')
set_option('use_real_price', True)
set_order_cost(OrderCost(open_tax=0, close_tax=0.0005,
    open_commission=0.0003, close_commission=0.0003,
    close_today_commission=0, min_commission=5), type='stock')
set_slippage(PriceRelatedSlippage(0.002))
run_daily(check_trailing_stop, time='open')
run_monthly(rebalance, 1, time='open+5m')
run_daily(record_risk_state, time='after_close')
        """,
    )

    add_heading(doc, "3 回测设计、参数调整与结果评价", 1)
    add_heading(doc, "3.1 数据范围、样本切分与撮合假设", 2)
    add_body(
        doc,
        f"独立复核使用 TASK3 已清洗的前复权日线数据，共{quality['security_count']}只股票、{quality['row_count']:,}条证券—交易日记录，数据覆盖{quality['calendar_start']}至{quality['calendar_end']}。2019年用于指标预热；2020—2023年作为开发期，2024—2025年作为验证期，2026年1月初至7月10日作为锁参观察期。参数只能使用开发期和验证期选择，观察期不参与选参。",
    )
    add_table_caption(doc, "表5 样本切分与评价职责")
    add_table(
        doc,
        ["阶段", "区间", "主要用途", "是否允许选参"],
        [
            ["预热期", "2019-01-02—2019-12-31", "形成最长200日指标窗口", "否"],
            ["开发期", "2020-01-01—2023-12-31", "排除明显不稳定组合", "是"],
            ["验证期", "2024-01-01—2025-12-31", "按夏普比率选择锁参方案", "是"],
            ["锁参观察期", "2026-01-01—2026-07-10", "模拟部署后的独立观察", "否"],
        ],
        [1600, 2450, 3375, 1600],
    )
    add_body(
        doc,
        "本地撮合在交易日开盘以昨日及更早数据形成信号，并以当日开盘价加减0.1%作为买卖价格；佣金为双边万分之三且最低5元，卖出另计万分之五印花税，股票按100股整手交易。基准为十股每日等权收益，用于控制股票池本身的市场表现。JoinQuant 页面中仍设置沪深300指数作为官方基准，因此本地数值与官方引擎不会逐点完全一致。",
    )
    add_heading(doc, "3.2 参数调整显著改善验证期表现", 2)
    add_table_caption(doc, "表6 基础模板与锁参策略的分期表现")
    add_table(
        doc,
        ["阶段/策略", "累计收益", "年化收益", "夏普", "最大回撤", "基准收益", "年化换手"],
        [
            ["验证期·基础", pct(baseline_validation["cumulative_return"]), pct(baseline_validation["annualized_return"]), num(baseline_validation["sharpe_ratio"]), pct(baseline_validation["max_drawdown"]), pct(baseline_validation["benchmark_return"]), f"{baseline_validation['annual_turnover']:.1f}倍"],
            ["验证期·锁参", pct(validation["cumulative_return"]), pct(validation["annualized_return"]), num(validation["sharpe_ratio"]), pct(validation["max_drawdown"]), pct(validation["benchmark_return"]), f"{validation['annual_turnover']:.1f}倍"],
            ["观察期·基础", pct(baseline_paper["cumulative_return"]), pct(baseline_paper["annualized_return"]), num(baseline_paper["sharpe_ratio"]), pct(baseline_paper["max_drawdown"]), pct(baseline_paper["benchmark_return"]), f"{baseline_paper['annual_turnover']:.1f}倍"],
            ["观察期·锁参", pct(paper["cumulative_return"]), pct(paper["annualized_return"]), num(paper["sharpe_ratio"]), pct(paper["max_drawdown"]), pct(paper["benchmark_return"]), f"{paper['annual_turnover']:.1f}倍"],
        ],
        [1900, 1225, 1225, 950, 1225, 1250, 1250],
    )
    add_figure(
        doc,
        "figure_2_nav_comparison.png",
        "图2 验证期与锁参观察期的累计净值",
        "图2显示，基础模板在2024年以后逐步落后，锁参策略的净值路径更稳定。锁参策略验证期收益为38.76%，仍低于十股等权基准的65.82%，说明风险控制降低回撤的同时放弃了部分上涨；进入2026年观察期后，策略短期跑赢基准，但该区间不足以证明结构性超额收益。",
    )
    add_figure(
        doc,
        "figure_3_period_returns.png",
        "图3 基础模板、锁参参数与基准的分期收益",
        "图3把参数调整的效果与基准机会成本放在同一视图。基础模板在验证期和观察期均为负收益；锁参策略两个阶段均为正。验证期仍明显落后基准，而观察期仅领先约2.06个百分点，因此结论应是“参数调整改善可行性”，而不是“策略已经稳定战胜市场”。",
    )
    add_heading(doc, "3.3 参数面并不平滑，仍需警惕过拟合", 2)
    add_body(
        doc,
        "参数搜索覆盖3个动量窗口、3个趋势窗口、3种持仓数量和3档止损阈值，共81组组合。开发期要求交易数量充分、夏普为正、最大回撤不超过55%；通过约束的组合再按验证期夏普比率排序，最大回撤和换手作为次级排序条件。最终选择20日动量、60日均线、3只持仓和8%止损。该规则不使用2026年数据，因此观察期仍保留独立性。",
    )
    add_figure(
        doc,
        "figure_4_parameter_heatmap.png",
        "图4 参数网格的验证期平均夏普比率",
        "图4对持仓数量和止损阈值取平均后展示动量—趋势窗口的参数面。20日动量整体优于60日和120日动量，但不同趋势窗口之间仍有差异，且平均值会掩盖止损和持仓数量的交互。参数面并非宽阔、平滑的高原，说明最终参数仍可能受到特定市场阶段影响。",
    )
    top_parameters = search.sort_values("validation_sharpe_ratio", ascending=False).head(6)
    add_table_caption(doc, "表7 验证期表现最高的六组参数")
    add_table(
        doc,
        ["参数", "开发期夏普", "验证期收益", "验证期夏普", "验证期回撤", "观察期收益"],
        [
            [
                row.parameter,
                num(row.development_sharpe_ratio),
                pct(row.validation_cumulative_return),
                num(row.validation_sharpe_ratio),
                pct(row.validation_max_drawdown),
                pct(row.paper_cumulative_return),
            ]
            for row in top_parameters.itertuples(index=False)
        ],
        [2300, 1350, 1350, 1350, 1350, 1325],
    )
    add_callout(
        doc,
        "稳健性判断",
        "部分验证期排名靠前的参数在2026年观察期转为负收益。该负面结果没有被删除，而是作为市场状态依赖和多重检验风险的直接证据。后续应使用滚动走步验证，而不是继续根据观察期结果重新挑参数。",
        fill=LIGHT_GRAY,
    )

    add_heading(doc, "4 模拟交易部署与运行推演", 1)
    add_heading(doc, "4.1 官方模拟交易部署步骤", 2)
    add_body(
        doc,
        "根据 JoinQuant 官方新手指引，完成完整回测后可以从“模拟交易”列表新建模拟交易，也可以从回测详情页直接创建。设置模拟交易参数时需要选择一个历史回测，平台会使用该回测对应的代码。创建成功后，至少要等待 A 股开盘一次才能看到模拟交易结果。因此，课程提交日前临时创建账户只能证明部署成功，不能形成有统计意义的长期业绩。",
    )
    add_bullet(doc, "在策略列表中新建策略，粘贴附带的 joinquant_strategy.py 代码并完成编译。")
    add_bullet(doc, "设置日频、初始资金100万元、回测区间2020-01-01至2025-12-31和沪深300基准，运行完整回测。")
    add_bullet(doc, "核对日志、订单、收益、最大回撤与 record 生成的仓位/现金比例曲线。")
    add_bullet(doc, "在回测详情页选择“新建模拟交易”，使用锁参代码，不再根据观察期盈亏调整参数。")
    add_bullet(doc, "等待至少一个A股交易日，检查策略是否按计划运行、订单是否被拒绝以及持仓是否符合目标。")
    add_heading(doc, "4.2 锁参版本、监控字段与运行检查", 2)
    add_table_caption(doc, "表8 模拟交易部署配置")
    add_table(
        doc,
        ["配置项", "锁定设置", "运行检查"],
        [
            ["策略版本", "M20/T60/N3/S8", "代码与回测版本一致"],
            ["初始资金", "1,000,000元", "现金和总资产初始化正确"],
            ["运行频率", "日频", "止损在开盘、调仓在月初执行"],
            ["基准", "沪深300指数", "结果页显示基准曲线"],
            ["成本", "佣金0.03%、卖出税0.05%、滑点0.20%", "订单费用与预期一致"],
            ["风险记录", "exposure、cash_ratio、holding_count", "收盘后记录连续且无缺口"],
            ["异常监控", "停牌、ST、涨跌停、订单为空", "日志能解释未成交原因"],
        ],
        [1850, 3400, 3775],
    )
    add_heading(doc, "4.3 锁参观察期取得正收益，但观察窗口仍短", 2)
    add_body(
        doc,
        f"本地锁参观察期从{paper['start_date']}至{paper['end_date']}，共{int(paper['observations'])}个交易日。策略累计收益{pct(paper['cumulative_return'])}，等权基准收益{pct(paper['benchmark_return'])}，超额收益{pct(paper['excess_return'], signed=True)}；夏普比率{num(paper['sharpe_ratio'])}，最大回撤{pct(paper['max_drawdown'])}。这是一段使用冻结参数的历史纸上推演，用于在官方模拟交易积累足够时间之前检查策略行为，不代表个人 JoinQuant 模拟账户已经运行相同天数。",
    )
    paper_trades = trades.loc[
        (pd.to_datetime(trades["trade_date"]) >= pd.Timestamp("2026-01-01"))
        & (pd.to_datetime(trades["trade_date"]) <= pd.Timestamp("2026-07-10"))
    ].copy()
    recent = paper_trades.tail(8)
    symbol_lookup = universe.set_index("ts_code")["stock_name"].to_dict()
    reason_map = {
        "trailing_stop": "跟踪止损",
        "rebalance_exit": "调仓退出",
        "rebalance_reduce": "调仓减仓",
        "rebalance_entry": "调仓建仓",
        "rebalance_add": "调仓加仓",
    }
    side_map = {"buy": "买入", "sell": "卖出"}
    add_table_caption(doc, "表9 锁参观察期最近八笔模拟订单")
    add_table(
        doc,
        ["日期", "证券", "方向", "数量", "成交金额", "原因"],
        [
            [
                row.trade_date,
                symbol_lookup.get(row.symbol, row.symbol),
                side_map.get(row.side, row.side),
                f"{int(row.quantity):,}",
                f"{row.gross_value:,.0f}元",
                reason_map.get(row.reason, row.reason),
            ]
            for row in recent.itertuples(index=False)
        ],
        [1600, 1500, 1000, 1150, 1900, 1875],
    )

    add_heading(doc, "5 风险暴露、稳健性与实盘差异", 1)
    add_heading(doc, "5.1 回撤改善不等于风险消失", 2)
    add_figure(
        doc,
        "figure_5_drawdown.png",
        "图5 验证期与锁参观察期的回撤路径",
        "图5表明锁参策略的验证期最大回撤约14.92%，明显小于基础模板的27.58%。不过，2026年观察期仍出现14.07%的最大回撤，接近验证期极值，说明较短动量和更紧止损并未消除市场急跌、跳空和相关性上升风险。模拟交易应对连续回撤设置预警，而不是等待止损完全解决组合风险。",
    )
    add_heading(doc, "5.2 交易成本与高换手削弱策略优势", 2)
    add_figure(
        doc,
        "figure_6_cost_sensitivity.png",
        "图6 验证期至锁参观察期的交易成本敏感性",
        f"图6显示，从2024年初至2026年7月10日，基准成本下策略累计收益为{pct(costs.loc[costs['multiplier'] == 1.0, 'cumulative_return'].iloc[0])}；成本加倍后降至{pct(costs.loc[costs['multiplier'] == 2.0, 'cumulative_return'].iloc[0])}。收益仍为正，但明显下滑。观察期年化换手约{paper['annual_turnover']:.1f}倍，说明真实冲击成本、最低佣金和未成交会进一步压缩结果。",
    )
    add_heading(doc, "5.3 行业与个股集中度仍是主要风险源", 2)
    add_figure(
        doc,
        "figure_7_industry_exposure.png",
        "图7 锁参观察期的行业风险暴露",
        f"图7按全部观察日计算行业平均权重，并同时展示单日最高权重。半导体行业单日最高权重达到{pct(industry_exposure['max_weight'].max())}，高于单股40%的目标上限，这是两只同业股票同时入选及价格漂移共同造成的。当前策略只限制单股，不限制行业，因此需要增加行业上限或相关性约束。",
    )
    add_table_caption(doc, "表10 锁参观察期个股暴露最高的五只股票")
    add_table(
        doc,
        ["证券", "行业", "全期平均权重", "持仓日平均权重", "单日最高权重", "持仓日数"],
        [
            [
                row.stock_name,
                row.industry,
                pct(row.average_weight),
                pct(row.active_day_average_weight),
                pct(row.max_weight),
                str(int(row.holding_days)),
            ]
            for row in stock_exposure.head(5).itertuples(index=False)
        ],
        [1500, 1400, 1600, 1750, 1600, 1175],
    )
    add_heading(doc, "5.4 尾部损失、模型边界与回测偏差", 2)
    add_figure(
        doc,
        "figure_8_tail_risk.png",
        "图8 锁参观察期的日收益分布与左尾风险",
        f"图8显示观察期日收益并非对称分布。历史95% VaR约为{pct(paper['daily_var_95'])}，进入最差5%交易日后的平均损失（CVaR）约为{pct(paper['daily_cvar_95'])}。这些数值只反映有限样本；涨跌停、跳空和多只持仓同时下跌时，实际尾部损失可能更大。",
    )
    add_table_caption(doc, "表11 主要风险暴露及应对措施")
    add_table(
        doc,
        ["风险", "当前证据", "实盘含义", "改进措施"],
        [
            ["市场方向", f"观察期β={paper['beta']:.2f}", "组合仍受大盘涨跌驱动", "设置组合波动率与回撤阈值"],
            ["个股集中", "目标上限40%，漂移后可超限", "单一公司事件影响净值", "增加再平衡容忍带和硬上限"],
            ["行业集中", f"行业单日最高{pct(industry_exposure['max_weight'].max())}", "同业相关性在压力期上升", "设置行业上限或相关性惩罚"],
            ["高换手", f"观察期年化{paper['annual_turnover']:.1f}倍", "成本和滑点侵蚀收益", "增加换仓阈值并降低无效替换"],
            ["尾部损失", f"95% CVaR {pct(paper['daily_cvar_95'])}", "跳空可能越过止损价", "组合级止损、现金缓冲与压力测试"],
            ["数据偏差", "固定当前股票池", "幸存者偏差高估可行性", "改用历史可交易股票池和点时数据"],
            ["引擎差异", "本地日频近似撮合", "与官方成交和订单状态不同", "以JoinQuant完整回测和模拟日志复核"],
        ],
        [1300, 2050, 2600, 3075],
    )

    add_heading(doc, "6 经验、教训与后续改进", 1)
    add_heading(doc, "6.1 平台实施过程中的经验", 2)
    add_body(
        doc,
        "第一，策略代码必须围绕平台的时间语义设计。日频历史函数不含当天数据，因此信号和成交时点要清楚区分。第二，参数优化的目标不是找到全样本最高收益，而是用开发期排除脆弱配置，再用验证期锁定参数，并让后续观察期保持独立。第三，回测结果必须同时比较基准、回撤、换手和暴露；正收益并不自动等于有效超额。第四，模拟交易的主要价值是发现平台执行问题，例如订单被拒绝、涨跌停无法成交、定时任务未触发和日志缺失，而不是在几天内证明收益能力。",
    )
    add_body(
        doc,
        "本次结果也给出一个重要反例：锁参策略在验证期和观察期均明显优于基础模板，但验证期仍落后等权基准。风险控制提高了可管理性，却牺牲部分牛市暴露。若只看策略自身收益，会高估其价值；若只追求基准超额，又可能忽略回撤和执行稳定性。实际部署需要在收益、风险和交易成本之间明确优先级。",
    )
    add_heading(doc, "6.2 从回测到模拟交易的提交前检查", 2)
    add_bullet(doc, "本人登录 JoinQuant，确认账号认证状态可用，报告中不展示任何敏感认证信息。")
    add_bullet(doc, "粘贴完整策略代码并编译，检查不存在未定义 API、缩进错误或 Python 版本问题。")
    add_bullet(doc, "运行锁参版本的完整回测，保存回测参数、结果页、风险指标和订单记录。")
    add_bullet(doc, "从完整回测创建模拟交易，确认初始资金、运行频率、策略版本和通知设置。")
    add_bullet(doc, "至少经历一个 A 股交易日，检查定时任务、持仓、未成交订单与 record 风险曲线。")
    add_bullet(doc, "若官方结果与本地复核差异较大，先核对复权、撮合时点、费用、停牌和涨跌停，不直接修改参数。")
    add_bullet(doc, "提交前确认文件名为“李子平TASK7.pdf”，图表编号连续、标题完整、文字可读。")

    add_heading(doc, "7 结论", 1)
    add_body(
        doc,
        "本文完成了从 JoinQuant 平台功能梳理、策略模板调整、参数验证、模拟部署设计到风险评估的完整流程。相较60日动量、120日趋势和12%止损的基础模板，20日动量、60日趋势和8%止损的锁参版本显著改善验证期收益和回撤，并在2026年锁参观察期取得9.65%的正收益和2.06个百分点的短期超额。然而，验证期仍落后等权基准，高换手、行业集中、固定股票池和本地撮合近似限制了结论强度。",
    )
    add_body(
        doc,
        "因此，策略可以进入 JoinQuant 模拟交易进行执行层验证，但尚不具备直接实盘投入的充分证据。下一阶段应在个人账户完成官方完整回测与模拟交易，积累至少三个月的订单、成交、持仓和风险记录；同时增加历史股票池、行业上限、换仓缓冲与滚动走步验证。只有当官方模拟结果在成本、回撤和执行稳定性上与研究结论一致时，才适合讨论更进一步的实盘部署。",
    )

    add_heading(doc, "参考文献", 1)
    references = [
        "[1] JoinQuant. 新手指引：注册、策略编辑、完整回测与模拟交易流程. https://www.joinquant.com/help/api/guide",
        "[2] JoinQuant. API文档：策略设置、历史数据、定时运行、下单与研究回测接口. https://cdn.joinquant.com/help/img/JoinQuantAPI.pdf",
        "[3] JoinQuant. 关于我们：平台数据、研究、回测与模拟交易能力说明. https://www.joinquant.com/about",
        "[4] 本地数据：PKU-WorkShop-202607/Task3/data，十只A股前复权日线，2019-01-02至2026-07-10。",
        "[5] 本地复核代码：TASK7/scripts/backtest_strategy.py；JoinQuant兼容代码：TASK7/scripts/joinquant_strategy.py。",
    ]
    for reference in references:
        add_body(doc, reference)

    add_heading(doc, "附录 核心策略代码与参数说明", 1)
    add_body(
        doc,
        "完整可运行策略见随报告交付的 joinquant_strategy.py。下列代码展示核心选股与目标持仓逻辑；函数使用昨日及更早数据形成信号，并在月初调用，避免直接使用当日收盘价交易。",
    )
    add_code_block(
        doc,
        """
closes = attribute_history(security, required, '1d',
    fields=['close'], skip_paused=True, df=True)['close']
latest = float(closes.iloc[-1])
momentum = latest / float(closes.iloc[-g.momentum_window - 1]) - 1.0
trend_average = float(closes.tail(g.trend_window).mean())
volatility = float(closes.pct_change().dropna()
    .tail(g.volatility_window).std() * np.sqrt(252))
if momentum > 0 and latest > trend_average and volatility > 0:
    candidate_momentum[security] = momentum
selected = sorted(candidate_momentum,
    key=candidate_momentum.get, reverse=True)[:g.top_n]
        """,
    )
    add_body(
        doc,
        "跟踪止损在每日开盘执行，比较昨日收盘价与持仓历史高点。当跌幅达到阈值时清仓，并将该证券加入当日禁止重新买入集合，避免止损和月度调仓在同一天相互抵消。",
    )
    add_code_block(
        doc,
        """
yesterday_high = float(history_data['high'].iloc[-1])
yesterday_close = float(history_data['close'].iloc[-1])
new_peak = max(g.high_watermark.get(security, yesterday_high),
               yesterday_high)
if yesterday_close <= new_peak * (1.0 - g.stop_loss):
    order_target_value(security, 0)
    g.stopped_today.add(security)
        """,
    )
    add_callout(
        doc,
        "平台复现说明",
        "在 JoinQuant 运行时，应以附带的完整 .py 文件为准。若平台 API 版本提示 avoid_future_data 不可用，代码中的 try/except 会跳过该选项；其余核心函数均按官方 API 文档设计。",
        fill=LIGHT_GRAY,
    )

    output = REPORT_DIR / f"{SUBMISSION_STEM}.docx"
    doc.save(output)
    return output


def write_supporting_notes(data: dict[str, object]) -> None:
    chart_rows = [
        ["图1", "平台流程", "策略如何从注册走到模拟交易", "流程图", "开发到模拟的五阶段流程", "figure_1_platform_workflow.png"],
        ["图2", "净值", "参数调整后路径是否改善", "趋势", "三条累计净值线", "figure_2_nav_comparison.png"],
        ["图3", "分期收益", "基础、锁参和基准在两个阶段如何比较", "比较", "分组柱形图", "figure_3_period_returns.png"],
        ["图4", "参数稳健性", "动量与趋势窗口是否存在宽阔高原", "矩阵", "带数值热力图", "figure_4_parameter_heatmap.png"],
        ["图5", "回撤", "参数调整是否改善损失路径", "趋势与基准", "多线回撤图", "figure_5_drawdown.png"],
        ["图6", "成本", "交易成本变化如何影响收益", "比较", "三情景柱形图", "figure_6_cost_sensitivity.png"],
        ["图7", "行业暴露", "行业平均与峰值集中度如何", "比较", "分组水平条形图", "figure_7_industry_exposure.png"],
        ["图8", "尾部风险", "观察期左尾损失有多大", "分布", "直方图与VaR/CVaR线", "figure_8_tail_risk.png"],
    ]
    pd.DataFrame(
        chart_rows,
        columns=["figure", "section", "question", "family", "chart_type", "artifact"],
    ).to_csv(OUTPUT_DIR / "chart_map.csv", index=False, encoding="utf-8-sig")
    notes = [
        "# TASK7 Report Source Notes",
        "",
        "- Audience: technical/coursework report.",
        "- Report shape: technical summary, evidence, scope and metrics, method, limitations, next steps.",
        "- Data: Task3/data ten-stock adjusted daily bars; provenance retained in TASK7/data/strategy_universe.csv.",
        "- Backtest: TASK7/scripts/backtest_strategy.py; JoinQuant implementation: TASK7/scripts/joinquant_strategy.py.",
        "- Official sources: JoinQuant beginner guide, API PDF and about page.",
        "- Benchmark note: local independent review uses the equal-weight return of the fixed ten-stock universe; JoinQuant deployment sets CSI 300 through set_benchmark.",
        "- Evidence boundary: no claim of completed identity verification, official backtest ID, simulation ID, or long-running JoinQuant paper account.",
        "- Document design: narrative_proposal preset with named TASK7 submission override: A4, SimSun 10.5 pt, 1.5 line spacing, 0 pt paragraph spacing, justified body; restrained editorial cover.",
        "- Omission: no screenshot of personal certification because it would expose sensitive account information.",
        "- Further question: whether official JoinQuant fills and costs materially change local replicated results.",
    ]
    (OUTPUT_DIR / "report_source_notes.md").write_text("\n".join(notes), encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--skip-figures", action="store_true")
    args = parser.parse_args()
    ensure_dirs()
    data = read_inputs()
    if not args.skip_figures:
        make_figures(data)
    write_supporting_notes(data)
    output = build_document(data)
    print(output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
