from __future__ import annotations

import json
import re
from pathlib import Path

import pdfplumber
from docx import Document
from docx.oxml.ns import qn
from PIL import Image


ROOT = Path(__file__).resolve().parents[2]
TASK7 = ROOT / "TASK7"
REPORT = TASK7 / "report"
OUTPUTS = TASK7 / "outputs"
FIGURES = OUTPUTS / "figures"
DOCX_PATH = REPORT / "李子平TASK7.docx"
PDF_PATH = REPORT / "李子平TASK7.pdf"
STRATEGY_PATH = TASK7 / "scripts" / "joinquant_strategy.py"


def check(condition: bool, message: str, errors: list[str]) -> None:
    if not condition:
        errors.append(message)


def main() -> int:
    errors: list[str] = []
    checks: list[str] = []

    check(DOCX_PATH.exists() and DOCX_PATH.stat().st_size > 200_000, "DOCX 不存在或文件过小", errors)
    check(PDF_PATH.exists() and PDF_PATH.stat().st_size > 200_000, "PDF 不存在或文件过小", errors)
    check(STRATEGY_PATH.exists() and STRATEGY_PATH.stat().st_size > 5_000, "JoinQuant 策略代码缺失", errors)
    if errors:
        print("\n".join(errors))
        return 1

    doc = Document(DOCX_PATH)
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    full_text = "\n".join(paragraphs)

    required = [
        "JoinQuant 策略部署与模拟交易实战",
        "摘要",
        "技术摘要",
        "1 JoinQuant 平台认知与操作流程",
        "2 策略模板调整与交易规则设计",
        "3 回测设计、参数调整与结果评价",
        "4 模拟交易部署与运行推演",
        "5 风险暴露、稳健性与实盘差异",
        "6 经验、教训与后续改进",
        "7 结论",
        "参考文献",
        "附录 核心策略代码与参数说明",
        "证据边界",
        "没有虚构回测",
    ]
    for item in required:
        check(item in full_text, f"缺少必需内容：{item}", errors)

    check("待更新" not in full_text and "[[TOC]]" not in full_text, "文档仍含占位符", errors)
    check("�" not in full_text, "文档正文含乱码替换字符", errors)

    abstract_start = full_text.find("摘要")
    abstract_end = full_text.find("关键词：", abstract_start)
    abstract_text = full_text[abstract_start + len("摘要") : abstract_end]
    abstract_han = re.findall(r"[\u4e00-\u9fff]", abstract_text)
    check(200 <= len(abstract_han) <= 400, f"摘要汉字数为 {len(abstract_han)}，不在 200–400 范围", errors)

    figure_numbers = [
        int(match.group(1))
        for text in paragraphs
        if (match := re.match(r"^图\s*(\d+)\s+", text))
    ]
    table_numbers = [
        int(match.group(1))
        for text in paragraphs
        if (match := re.match(r"^表\s*(\d+)\s+", text))
    ]
    check(sorted(set(figure_numbers)) == list(range(1, 9)), f"图号不连续：{figure_numbers}", errors)
    check(sorted(set(table_numbers)) == list(range(1, 12)), f"表号不连续：{table_numbers}", errors)

    inline_count = len(doc.element.body.xpath(".//wp:inline"))
    check(inline_count == 8, f"嵌入图片数量应为 8，实际为 {inline_count}", errors)

    normal = doc.styles["Normal"]
    ascii_font = normal.font.name
    east_asia_font = normal.element.rPr.rFonts.get(qn("w:eastAsia"))
    size_pt = normal.font.size.pt if normal.font.size else None
    check(ascii_font == "SimSun" and east_asia_font == "SimSun", f"正文样式字体不为宋体：{ascii_font}/{east_asia_font}", errors)
    check(size_pt is not None and abs(size_pt - 10.5) < 0.05, f"正文样式字号不是 10.5 磅：{size_pt}", errors)
    check(normal.paragraph_format.line_spacing == 1.5, f"正文样式行距不是 1.5：{normal.paragraph_format.line_spacing}", errors)
    before = normal.paragraph_format.space_before
    after = normal.paragraph_format.space_after
    check(before is not None and before.pt == 0, f"正文段前距不是 0：{before}", errors)
    check(after is not None and after.pt == 0, f"正文段后距不是 0：{after}", errors)

    for index, section in enumerate(doc.sections, start=1):
        width = section.page_width.cm
        height = section.page_height.cm
        check(abs(width - 21.0) < 0.1 and abs(height - 29.7) < 0.1, f"第 {index} 节不是 A4：{width:.2f}×{height:.2f} cm", errors)

    toc_table = next(
        (
            table
            for table in doc.tables
            if any(row.cells[0].text.strip() == "摘要" for row in table.rows)
        ),
        None,
    )
    check(toc_table is not None, "未找到静态目录表", errors)
    if toc_table is not None:
        for row in toc_table.rows:
            title = row.cells[0].text.strip()
            page = row.cells[1].text.strip()
            check(bool(re.fullmatch(r"\d+", page)), f"目录页码无效：{title} -> {page}", errors)

    with pdfplumber.open(PDF_PATH) as pdf:
        check(len(pdf.pages) == 17, f"PDF 页数不是 17：{len(pdf.pages)}", errors)
        extracted_pages = []
        for page_number, page in enumerate(pdf.pages, start=1):
            check(abs(float(page.width) - 595.3) < 2 and abs(float(page.height) - 841.9) < 2, f"PDF 第 {page_number} 页不是 A4", errors)
            text = page.extract_text() or ""
            extracted_pages.append(text)
            check(len(text.strip()) > 20, f"PDF 第 {page_number} 页文本过少，可能为空白", errors)
        pdf_text = "\n".join(extracted_pages)
        check("李子平" in pdf_text and "TASK7" in pdf_text, "PDF 未提取到作者或 TASK7", errors)
        check("�" not in pdf_text, "PDF 提取文本含乱码替换字符", errors)

    figure_files = sorted(FIGURES.glob("figure_*.png"))
    check(len(figure_files) == 8, f"图文件数量应为 8，实际为 {len(figure_files)}", errors)
    for figure in figure_files:
        with Image.open(figure) as image:
            check(image.width >= 1200 and image.height >= 650, f"图分辨率偏低：{figure.name} {image.size}", errors)

    expected_outputs = [
        "strategy_comparison.csv",
        "parameter_search.csv",
        "daily_backtest.csv",
        "trade_log.csv",
        "daily_holdings.csv",
        "cost_sensitivity.csv",
        "stock_exposure.csv",
        "industry_exposure.csv",
        "selected_parameters.json",
        "data_quality.json",
        "toc_pages.json",
        "report_source_notes.md",
    ]
    for name in expected_outputs:
        check((OUTPUTS / name).exists(), f"缺少分析输出：{name}", errors)

    strategy_text = STRATEGY_PATH.read_text(encoding="utf-8")
    for token in [
        "def initialize",
        "set_benchmark",
        "set_order_cost",
        "set_slippage",
        "run_monthly",
        "run_daily",
        "attribute_history",
        "order_target_value",
        "record",
    ]:
        check(token in strategy_text, f"JoinQuant 代码缺少：{token}", errors)

    checks.extend(
        [
            "DOCX/PDF 文件完整性",
            "章节、摘要、证据边界与无占位符",
            "图 1–8、表 1–11 连续编号",
            "宋体五号、1.5 倍行距、0 段间距、A4",
            "目录实际页码",
            "PDF 17 页文本与版面",
            "8 幅图的分辨率",
            "分析输出与 JoinQuant API 结构",
        ]
    )
    result = {
        "status": "PASS" if not errors else "FAIL",
        "checks": checks,
        "errors": errors,
        "docx_bytes": DOCX_PATH.stat().st_size,
        "pdf_bytes": PDF_PATH.stat().st_size,
        "pdf_pages": 17,
        "figures": len(figure_files),
        "tables": len(set(table_numbers)),
    }
    output_path = OUTPUTS / "validation_report.json"
    output_path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")

    if errors:
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 1
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
